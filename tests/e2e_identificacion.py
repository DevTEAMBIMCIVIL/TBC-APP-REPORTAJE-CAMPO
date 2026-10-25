"""Pruebas de la v3.1 y la v3.1.1 en un Pixel 7 emulado: cabecera en gris corporativo con texto negro, orden de la
identificación (expediente primero), técnico, cargo y objeto de listas desplegables (con «Otro» y detalle), cargo de
los asistentes, sugerencias de expedientes, técnico y cargo recordados en el teléfono, revisión, PDF con el formato
«Acta de Campo» (cabecera y pie en gris, v3.1.1), texto compartido y registro en Excel, migración de las actas de la
v3.0 (esquema 3), actas emitidas con la v3.0 y la v3.1 que conservan el formato con el que se emitieron, e importación.

Uso (desde la raíz del repo, con los mismos requisitos que e2e_seguimiento.py):
    python tests/e2e_identificacion.py
"""
import hashlib
import json
import pathlib
import re
import sys
import zipfile

from openpyxl import load_workbook
from playwright.sync_api import sync_playwright
from pypdf import PdfReader

sys.path.insert(0, str(pathlib.Path(__file__).parent))
from comun import (MARCO_GRIS, MARCO_NEGRO, RAIZ, SALIDA, Resultados, abrir, abrir_mis_actas,  # noqa: E402
                   cerrar_aviso, contexto_movil, esperar_fotos, firmar, marco_pdf, servidor, texto_pdf, vigilar_red)
from comun import anadir_asistente, escribir_empresa, escribir_expediente, escribir_responsable  # noqa: E402
from comun import avanzar_hasta, identificar  # noqa: E402
from generar_fotos import generar  # noqa: E402

TECNICOS = ["Alberto de Álava", "Ana Cascajosa", "Ángel García", "Beatriz Sánchez", "Blas González", "Carlos Rubio",
            "Francisco García", "Francisco S. Ortiz", "Javier Gómez", "José A. Gómez", "José L. Gutiérrez", "María Quinta",
            "Olga Provincial", "Sebastián Gil"]                  # en orden alfabético (14 desde la v3.1.3)
CARGOS = ["Inspector de Campo", "Redactor de Estudio", "Director de Obra", "Supervisor de Obra"]
OBJETOS = ["Licitación de Obra", "Inspección de Tajos de Obra", "Inventario", "Estudio de Patología", "Otro"]
EXP = "EXP-ID"
GRIS = "rgb(136, 136, 136)"                                     # #888888
DATOS_V312 = pathlib.Path(__file__).parent / "datos" / "v312"     # actas exportadas con la v3.1.2 (ver su LEEME.md)
DATOS_V313 = pathlib.Path(__file__).parent / "datos" / "v313"     # actas emitidas y exportadas con la v3.1.3 (ver su LEEME.md)
# Fase de Acuerdos en la página: la 5 en la v3.1.x y FASE.ACUERDOS (6) desde la v3.2, con la fase Comprobaciones.
ACUERDOS = "(typeof FASE !== 'undefined' ? FASE.ACUERDOS : 5)"
# Actas emitidas para inyectar en el almacén de la app: [{ id, num, expediente, momento, asistentes, indicaciones, … }]
INYECTAR = """async actas => {
  for (const x of actas) {
    const m = x.momento, emitida = (x.estado || 'emitida') === 'emitida', rev = x.revision || 0;
    await Almacen.guardarActa({ id: x.id, esquema: x.esquema || 4, estado: emitida ? 'emitida' : 'borrador', creada: m, modificada: m,
      emitida: emitida ? m : null, ...(emitida ? { versionEmision: x.version || '3.1.2' } : {}), revision: rev, fase: x.fase || 8,
      corrigeA: x.corrigeA || null, sustituidaPor: x.sustituidaPor || null,
      revisiones: Array.from({ length: rev + 1 }, (_, n) => ({ rev: n, fecha: emitida || n < rev ? m : null, motivo: n ? 'Corrección' : 'Emisión inicial' })),
      datos: { actaNum: x.num, numAuto: false, tecnico: x.tecnico || 'Blas González', cargoTecnico: 'Director de Obra',
               expediente: x.expediente, objeto: 'tajos', descripcion: '', fecha: m.slice(0, 10), hora: '09:00' },
      gps: { lat: '37.389210', lng: '-5.984460', acc: 8, hora: m }, sinGps: { activo: false, referencia: '', lat: '', lng: '' },
      condiciones: { meteo: [], temperatura: '' },
      asistentes: (x.asistentes || []).map((a, i) => ({ id: `${x.id}-as${i}`, nombre: a[0], empresa: a[1], cargo: a[2], firma: true })),
      observaciones: '', seguimiento: [], seguimientoOrigen: null, seguimientoNotas: [], seguimientoGuardado: null,
      indicaciones: (x.indicaciones || []).map((t, i) => ({ id: `${x.id}-ind${i}`, texto: t[0], responsable: t[1], plazo: '' })),
      fotos: [], firmasAsistentes: {}, firma: [{ refW: 300, refH: 150, puntos: [[10, 10], [60, 40], [120, 20], [180, 60]] }] });
  }
}"""
GRUPOS = """(() => { const s = document.getElementById('anadirAsistenteLista');
  return { sueltas: Array.from(s.children).filter(c => c.tagName === 'OPTION').map(o => o.textContent),
           grupos: Array.from(s.querySelectorAll('optgroup')).map(g => [g.label, Array.from(g.children).map(o => o.textContent)]) }; })()"""
DAÑOS = "Revisión de daños tras la tormenta"
LOSA = "Hormigonado de la losa de planta 1"
OPCIONES = "id => Array.from(document.getElementById(id).options).map(o => [o.value, o.textContent])"
ESTILOS = """(() => { const c = s => getComputedStyle(document.querySelector(s));
  return { cabecera: c('.app-header').backgroundColor, pasos: c('.stepper').backgroundColor,
           titulo: c('.header-text h1').color, sombra: c('.header-text h1').textShadow,
           subtitulo: c('.header-text p').color, guardado: c('.estado-guardado').color,
           nombre: document.querySelector('.header-text p').textContent, pestana: document.title,
           paso: c('.step-item:not(.active) .step-label').color, activo: c('.step-item.active .step-label').color,
           tema: document.querySelector('meta[name=theme-color]').content }; })()"""


def rgb(texto):
    return [int(v) for v in re.findall(r"\d+", texto)[:3]]


def contraste(a, b):
    """Relación de contraste WCAG entre dos colores CSS «rgb(…)»."""
    def lum(color):
        canales = [c / 255 for c in rgb(color)]
        lin = [c / 12.92 if c <= 0.03928 else ((c + 0.055) / 1.055) ** 2.4 for c in canales]
        return 0.2126 * lin[0] + 0.7152 * lin[1] + 0.0722 * lin[2]
    claro, oscuro = sorted((lum(a), lum(b)), reverse=True)
    return (claro + 0.05) / (oscuro + 0.05)


def descargar(pg, selector, ruta):
    with pg.expect_download(timeout=60000) as d:
        pg.click(selector)
    d.value.save_as(str(ruta))
    return d.value.suggested_filename


def zip_de_acta(ruta, acta):
    """Un ZIP como los de «Exportar», con el acta dada (sin fotos)."""
    with zipfile.ZipFile(ruta, "w") as z:
        z.writestr(f"{acta['datos']['actaNum']}_Rev{acta['revision']}/acta.json",
                   json.dumps({"formato": "tbc-acta", "version": 1, "app": "3.1.3", "acta": acta}, ensure_ascii=False))
    return ruta


def huellas(ruta):
    """SHA-256 del contenido de cada página del PDF (texto, trazos y colocación de imágenes)."""
    return [hashlib.sha256(pagina.get_contents().get_data()).hexdigest() for pagina in PdfReader(str(ruta)).pages]


def opciones(pg, id_select):
    return [o[0] for o in pg.evaluate(OPCIONES, id_select)]


def importar(pg, ruta):
    abrir_mis_actas(pg)
    pg.set_input_files("#importInput", str(ruta))
    pg.wait_for_selector("#alertOverlay.open", timeout=60000)
    aviso = cerrar_aviso(pg)
    pg.evaluate("cerrarMisActas()")
    return aviso


def revision_v313(p, navegador, r, externas, errores, foto):
    """Casos de la revisión independiente de la v3.1.3 (R1 a R8). Cada caso va en su propio teléfono y un caso que no
    llega a su comprobación (una espera agotada) cuenta como fallo sin detener los demás."""
    abiertos = []

    def telefono(actas=None, recordados=None):
        ctx = contexto_movil(p, navegador)
        abiertos.append(ctx)
        vigilar_red(ctx, externas)
        pg = abrir(ctx, errores)
        if actas or recordados:
            if actas:
                pg.evaluate(INYECTAR, actas)
            if recordados:
                pg.evaluate("r => localStorage.setItem('tbc-recordados', JSON.stringify(r))", recordados)
            pg.reload()
            pg.wait_for_function("state.acta && !document.getElementById('expedienteLista').classList.contains('hidden')", timeout=10000)
        return ctx, pg

    def caso(nombre, funcion):
        try:
            funcion()
        except Exception as e:  # noqa: BLE001 · una espera agotada es un fallo de esta comprobación, no de la batería
            r.comprobar(nombre, False, f"{type(e).__name__}: {str(e).splitlines()[0][:140]}")
        finally:
            while abiertos:
                abiertos.pop().close()

    def descargar_y_emitir(pg, ruta):
        pg.wait_for_function("document.getElementById('emitPdfEstado').textContent.startsWith('PDF listo')", timeout=60000)
        descargar(pg, "#btnDescargar", ruta)
        pg.wait_for_function("state.acta.estado === 'emitida'")

    una = [{"id": "base-1", "num": "EXP-BASE-AV-01", "expediente": "EXP-BASE", "momento": "2026-10-01T09:00:00.000Z"}]

    # ── R1 · Tras importar o restaurar una copia, y cuando otra ventana guarda, los desplegables se ponen al día ──
    R1 = "Tras importar o restaurar una copia, y si otra ventana guarda un acta, el desplegable y los asistentes ofrecen lo nuevo"

    def r1():
        ctx, pg = telefono()
        sin_lista = pg.is_hidden("#expedienteLista")
        importar(pg, DATOS_V312 / "av01.zip")
        pg.wait_for_function("Array.from(document.getElementById('expedienteLista').options).some(o => o.value === 'EXP-V312')",
                             timeout=5000)
        tras_importar = [pg.is_visible("#expedienteLista"), opciones(pg, "expedienteLista")]
        pg.select_option("#tecnico", "Blas González")
        pg.evaluate("irAFase(2)")
        pg.wait_for_function("""Array.from(document.getElementById('anadirAsistenteLista').options)
            .some(o => o.textContent.startsWith('Juan Pérez Gómez'))""", timeout=5000)
        pg.evaluate("irAFase(1)")
        otra = abrir(ctx, errores)                           # la misma app en otra pestaña
        otra.evaluate(INYECTAR, [{"id": "r1-otra", "num": "EXP-OTRA-AV-01", "expediente": "EXP-OTRA", "momento": "2026-10-30T09:00:00.000Z"}])
        otra.evaluate("avisarVentanas({ tipo: 'guardada', id: 'r1-otra' })")
        pg.wait_for_function("Array.from(document.getElementById('expedienteLista').options).some(o => o.value === 'EXP-OTRA')",
                             timeout=5000)
        cerrar_aviso(pg)
        r.comprobar(R1, sin_lista and tras_importar == [True, ["", "EXP-V312", "__nuevo"]]
                    and opciones(pg, "expedienteLista") == ["", "EXP-OTRA", "EXP-V312", "__nuevo"], str(tras_importar))
        otra.close()
    caso(R1, r1)

    # ── R7 · Tocar el desplegable sin cambiar el expediente no guarda un acta vacía ──
    R7 = "Elegir «Nuevo expediente…» en un acta nueva sin escribir nada no la guarda (no queda un borrador «Sin número»)"

    def r7():
        ctx, pg = telefono(una)
        antes = pg.evaluate("Almacen.listarActas().then(l => l.length)")
        pg.select_option("#expedienteLista", "__nuevo")
        pg.wait_for_timeout(800)
        r.comprobar(R7, pg.evaluate("state.persistida") is False and pg.evaluate("Almacen.listarActas().then(l => l.length)") == antes
                    and pg.evaluate("document.activeElement.id") == "expediente")
    caso(R7, r7)

    # ── R8 · Con las flechas «Añadir asistente…» no añade (Intro, sí); nombre accesible del desplegable del expediente ──
    R8 = ("Teclado: las flechas recorren «Añadir asistente…» sin añadir nadie, Intro o salir añaden el elegido y Escape lo deja; "
          "el desplegable del expediente tiene nombre")

    def r8():
        ctx, pg = telefono(una)
        pg.select_option("#expedienteLista", "EXP-BASE")
        pg.select_option("#tecnico", "Blas González")
        pg.evaluate("irAFase(2)")
        pg.focus("#anadirAsistenteLista")
        for _ in range(3):
            pg.keyboard.press("ArrowDown")
        con_flechas = pg.evaluate("state.acta.asistentes.length")
        pg.keyboard.press("Enter")
        tras_intro = pg.evaluate("state.acta.asistentes.map(a => [a.nombre, a.empresa])")
        pg.focus("#anadirAsistenteLista")
        pg.keyboard.press("ArrowDown")
        pg.keyboard.press("ArrowDown")
        pg.keyboard.press("Escape")
        tras_escape = pg.evaluate("[state.acta.asistentes.length, document.getElementById('anadirAsistenteLista').value]")
        pg.focus("#anadirAsistenteLista")
        pg.keyboard.press("ArrowDown")
        pg.keyboard.press("ArrowDown")
        pg.click("#tecnicoNombreAsist")                       # sale del desplegable: se añade lo elegido
        tras_salir = pg.evaluate("state.acta.asistentes.length")
        r.comprobar(R8, con_flechas == 0 and len(tras_intro) == 1 and tras_intro[0][1] == "TEAMBIMCIVIL S.L."
                    and tras_escape == [1, ""] and tras_salir == 2
                    and pg.evaluate("document.getElementById('expedienteLista').getAttribute('aria-label')") == "Expediente del trabajo",
                    f"{con_flechas} {tras_intro} {tras_escape} {tras_salir}")
    caso(R8, r8)

    # ── R2 · Empresa y responsable: lo escrito se ofrece en la ficha siguiente, y «Otra…» parte de lo escrito ──
    R2 = "Empresa: la escrita en una ficha se ofrece en la siguiente y «Otra…» parte de ella para corregirla; también el responsable"

    def r2():
        ctx, pg = telefono(una)
        pg.select_option("#expedienteLista", "EXP-BASE")
        pg.select_option("#tecnico", "Blas González")
        pg.evaluate("irAFase(2)")
        anadir_asistente(pg)
        pg.fill("#asNombre0", "Juan Ruiz")
        escribir_empresa(pg, 0, "Constructora Surr")
        anadir_asistente(pg)
        en_siguiente = "Constructora Surr" in opciones(pg, "asEmpresaLista1")
        pg.click("[data-asistente='1'] .quitar-btn")
        pg.evaluate("irAFase(3)")
        pg.evaluate("irAFase(2)")
        vuelta = pg.evaluate("[document.getElementById('asEmpresaLista0').value, state.acta.asistentes[0].empresa]")
        pg.select_option("#asEmpresaLista0", "__otro")
        corregir = pg.evaluate("[document.getElementById('asEmpresa0').value, state.acta.asistentes[0].empresa, document.activeElement.id]")
        pg.fill("#asEmpresa0", "Constructora Sur")
        pg.evaluate("irAFase(5)")
        pg.click("button:has-text('Añadir indicación')")
        escribir_responsable(pg, 0, "Constructora Sur")
        pg.select_option("#indRespLista0", "__otro")
        responsable = pg.evaluate("[document.getElementById('indResp0').value, state.acta.indicaciones[0].responsable]")
        r.comprobar(R2, en_siguiente and vuelta == ["Constructora Surr", "Constructora Surr"]
                    and corregir == ["Constructora Surr", "Constructora Surr", "asEmpresa0"]
                    and pg.evaluate("state.acta.asistentes[0].empresa") == "Constructora Sur"
                    and responsable == ["Constructora Sur", "Constructora Sur"], f"{vuelta} {corregir} {responsable}")
    caso(R2, r2)

    acta_ch = {"id": "r3-a1", "num": "EXP-CH-AV-01", "expediente": "EXP-CH", "momento": "2026-10-01T09:00:00.000Z",
               "asistentes": [["Ana Cascajosa", "TEAMBIMCIVIL", "Redactor de Estudio"], ["Blas González", "TEAMBIMCIVIL", "Director de Obra"],
                              ["Juan Pérez", "Constructora Sur", "Jefe de obra"]],
               "indicaciones": [["Retirar el acopio", "TEAMBIMCIVIL · Director de Obra"], ["Limpiar el vial", "Constructora Sur · Jefe de obra"]]}

    # ── R3 · Los chips de visitas anteriores no repiten a quien ya está ni proponen al técnico ──
    R3 = "Chips de visitas anteriores: sin el técnico del acta ni quien ya se añadió del desplegable (TEAMBIMCIVIL o S.L.)"

    def r3():
        ctx, pg = telefono([acta_ch])
        pg.select_option("#expedienteLista", "EXP-CH")
        pg.select_option("#tecnico", "Blas González")
        pg.evaluate("irAFase(2)")
        pg.wait_for_selector("#asistentesSugeridos:not(.hidden)", timeout=5000)
        chips = pg.evaluate("Array.from(document.querySelectorAll('#asistentesSugeridos button')).map(b => b.textContent)")
        anadir_asistente(pg, "Ana Cascajosa · TEAMBIMCIVIL S.L. · Redactor de Estudio")
        pg.wait_for_timeout(500)
        chips2 = pg.evaluate("Array.from(document.querySelectorAll('#asistentesSugeridos button')).map(b => b.textContent)")
        r.comprobar(R3, chips == ["+ Ana Cascajosa · TEAMBIMCIVIL · Redactor de Estudio", "+ Juan Pérez · Constructora Sur · Jefe de obra"]
                    and chips2 == ["+ Juan Pérez · Constructora Sur · Jefe de obra"], f"{chips} {chips2}")
    caso(R3, r3)

    # ── R4 · «Nueva visita» y «Corregir» de un acta de la v3.1.2: los de TEAMBIMCIVIL, con la razón social ──
    R4 = "«Nueva visita» y «Corregir» de un acta de la v3.1.2: los de TEAMBIMCIVIL pasan a TEAMBIMCIVIL S.L. (la emitida no cambia)"

    def r4():
        ctx, pg = telefono([acta_ch])
        pg.evaluate("nuevaVisita('r3-a1')")
        pg.wait_for_function("state.acta.id !== 'r3-a1' && state.acta.asistentes.length === 3", timeout=10000)
        visita = pg.evaluate("state.acta.asistentes.map(a => a.empresa)")
        pg.evaluate("pedirTexto = async () => 'Corrección de prueba'")
        pg.evaluate("corregirActa('r3-a1')")
        pg.wait_for_function("state.acta.revision === 1 && state.acta.corrigeA === 'r3-a1'", timeout=10000)
        cerrar_aviso(pg)
        revision = pg.evaluate("[state.acta.asistentes.map(a => a.empresa), state.acta.indicaciones.map(i => i.responsable)]")
        original = pg.evaluate("Almacen.leerActa('r3-a1').then(a => [a.asistentes.map(x => x.empresa), a.indicaciones.map(i => i.responsable)])")
        r.comprobar(R4, visita == ["TEAMBIMCIVIL S.L.", "TEAMBIMCIVIL S.L.", "Constructora Sur"]
                    and revision == [["TEAMBIMCIVIL S.L.", "TEAMBIMCIVIL S.L.", "Constructora Sur"],
                                     ["TEAMBIMCIVIL S.L. · Director de Obra", "Constructora Sur · Jefe de obra"]]
                    and original == [["TEAMBIMCIVIL", "TEAMBIMCIVIL", "Constructora Sur"],
                                     ["TEAMBIMCIVIL · Director de Obra", "Constructora Sur · Jefe de obra"]], f"{visita} {revision}")
    caso(R4, r4)

    # ── R5 · El desplegable muestra el expediente como lo guarda el acta; las grafías con otros espacios, por separado ──
    R5 = "El desplegable muestra el expediente tal como lo guarda el acta, y las grafías con distintos espacios van por separado"

    def r5():
        ctx, pg = telefono([
            {"id": "r5-b1", "num": "Exp-B-AV-01", "expediente": "Exp-B", "momento": "2026-10-02T09:00:00.000Z"},
            {"id": "r5-b2", "num": "EXP-B-AV-02", "expediente": "EXP-B", "momento": "2026-10-09T09:00:00.000Z"},
            {"id": "r5-c1", "num": "OBRA  NORTE-AV-01", "expediente": "OBRA  NORTE", "momento": "2026-10-03T09:00:00.000Z"},
            {"id": "r5-c2", "num": "OBRA NORTE-AV-02", "expediente": "OBRA NORTE", "momento": "2026-10-04T09:00:00.000Z"}])
        grafias = opciones(pg, "expedienteLista")
        pg.evaluate("nuevaVisita('r5-b1')")
        pg.wait_for_function("state.acta.id !== 'r5-b1' && state.acta.datos.actaNum !== ''", timeout=10000)
        pg.wait_for_function("document.getElementById('expedienteLista').value !== ''", timeout=5000)
        exacta = pg.evaluate("[document.getElementById('expedienteLista').value, state.acta.datos.expediente, state.acta.datos.actaNum]")
        r.comprobar(R5, exacta == ["Exp-B", "Exp-B", "Exp-B-AV-03"] and grafias == ["", "EXP-B", "OBRA NORTE", "OBRA  NORTE", "__nuevo"],
                    f"{exacta} {grafias}")
    caso(R5, r5)

    # ── R6 · Un expediente mal escrito y corregido con una revisión deja de proponerse; lo de otras actas, no ──
    R6 = "Un expediente mal escrito y corregido con una revisión deja de proponerse (también tras borrar la Rev. 0); el recordado de otras actas sigue"

    def r6():
        ctx, pg = telefono(recordados={"expedientes": [{"valor": "EXP-LEGIT", "fecha": "2026-09-01T09:00:00.000Z",
                                                        "actas": ["acta-ya-borrada"]}], "personas": []})
        cdp = ctx.new_cdp_session(pg)
        pg.select_option("#expedienteLista", "__nuevo")
        identificar(pg, expediente="425_2026.COM")          # errata: COM por CON
        pg.click("#gpsBtn")
        pg.wait_for_selector("#gpsStatus.ok", timeout=15000)
        pg.wait_for_function("state.acta.datos.actaNum === '425_2026.COM-AV-01'", timeout=5000)
        avanzar_hasta(pg, 3)
        pg.set_input_files("#photoInput", [foto])
        esperar_fotos(pg, 1)
        avanzar_hasta(pg, 6)
        firmar(cdp, pg)
        avanzar_hasta(pg, 8)
        descargar_y_emitir(pg, SALIDA / "S313-rev0.pdf")
        rev0 = pg.evaluate("state.acta.id")
        pg.evaluate("pedirTexto = async () => 'Corrección del expediente'")
        pg.evaluate(f"corregirActa('{rev0}')")
        pg.wait_for_function("state.acta.revision === 1", timeout=10000)
        cerrar_aviso(pg)
        pg.evaluate("irAFase(1)")
        pg.select_option("#expedienteLista", "__nuevo")
        pg.fill("#expediente", "425_2026.CON")
        avanzar_hasta(pg, 6)
        firmar(cdp, pg)
        avanzar_hasta(pg, 8)
        descargar_y_emitir(pg, SALIDA / "S313-rev1.pdf")
        pg.evaluate("nuevaActa()")
        pg.wait_for_function("state.acta && !state.persistida && !document.getElementById('expedienteLista').classList.contains('hidden')")
        pg.wait_for_timeout(500)
        con_rev0 = opciones(pg, "expedienteLista")
        pg.evaluate(f"Almacen.borrarActa('{rev0}')")
        pg.evaluate("cerrarMisActas()")
        pg.wait_for_timeout(500)
        sin_rev0 = opciones(pg, "expedienteLista")
        recordados = pg.evaluate("JSON.parse(localStorage.getItem('tbc-recordados')).expedientes.map(e => e.valor)")
        r.comprobar(R6, con_rev0 == sin_rev0 == ["", "425_2026.CON", "EXP-LEGIT", "__nuevo"] and recordados == ["425_2026.CON", "EXP-LEGIT"],
                    f"{con_rev0} {sin_rev0} {recordados}")
    caso(R6, r6)

    # ── S1 · Los chips tampoco proponen a la gente de un acta corregida ──
    S1 = "Chips: tampoco proponen a la gente de un acta corregida (sustituida, o la que corrige la revisión en pantalla)"

    def s1():
        ctx, pg = telefono([
            {"id": "s1-r0", "num": "EXP-S1-AV-01", "expediente": "EXP-S1", "momento": "2026-10-01T09:00:00.000Z",
             "asistentes": [["Jaun Pérez", "Constructora Sur", "Jefe de obra"]], "sustituidaPor": "s1-r1"},
            {"id": "s1-r1", "num": "EXP-S1-AV-01", "revision": 1, "corrigeA": "s1-r0", "expediente": "EXP-S1",
             "momento": "2026-10-02T09:00:00.000Z", "asistentes": [["Juan Pérez", "Constructora Sur", "Jefe de obra"]]},
            {"id": "s1-b0", "num": "EXP-S1B-AV-01", "expediente": "EXP-S1B", "momento": "2026-10-03T09:00:00.000Z",
             "asistentes": [["Pepe Errata", "Constructora Sur", ""]]}])
        pg.select_option("#expedienteLista", "EXP-S1")
        pg.select_option("#tecnico", "Blas González")
        pg.evaluate("irAFase(2)")
        pg.wait_for_selector("#asistentesSugeridos:not(.hidden)", timeout=5000)
        chips = pg.evaluate("Array.from(document.querySelectorAll('#asistentesSugeridos button')).map(b => b.textContent)")
        pg.evaluate("pedirTexto = async () => 'Corrección del asistente'")
        pg.evaluate("corregirActa('s1-b0')")
        pg.wait_for_function("state.acta.revision === 1 && state.acta.corrigeA === 's1-b0'", timeout=10000)
        cerrar_aviso(pg)
        pg.evaluate("irAFase(2)")
        pg.click("[data-asistente='0'] .quitar-btn")
        pg.wait_for_timeout(500)
        chips2 = pg.evaluate("Array.from(document.querySelectorAll('#asistentesSugeridos button')).map(b => b.textContent)")
        r.comprobar(S1, chips == ["+ Juan Pérez · Constructora Sur · Jefe de obra"] and chips2 == [], f"{chips} {chips2}")
    caso(S1, s1)

    # ── S2 · Un borrador empezado con la v3.1.2 lleva lo de TEAMBIMCIVIL con la razón social, como su técnico ──
    S2 = "Un borrador empezado con la v3.1.2: sus asistentes y responsables de TEAMBIMCIVIL pasan a TEAMBIMCIVIL S.L., como el técnico"

    def s2():
        ctx, pg = telefono([{"id": "s2-b", "num": "EXP-S2-AV-01", "expediente": "EXP-S2", "momento": "2026-10-04T09:00:00.000Z",
                             "estado": "borrador", "fase": 2, "asistentes": [["Ana Cascajosa", "TEAMBIMCIVIL", "Redactor de Estudio"]],
                             "indicaciones": [["Retirar el acopio", "TEAMBIMCIVIL · Director de Obra"]]}])
        pg.evaluate("abrirActaGuardada('s2-b')")
        pg.wait_for_function("state.acta.id === 's2-b'")
        datos = pg.evaluate("[state.acta.asistentes[0].empresa, state.acta.indicaciones[0].responsable]")
        pg.evaluate(f"irAFase({ACUERDOS})")
        en_lista = pg.evaluate("document.getElementById('indRespLista0').value")
        r.comprobar(S2, datos == ["TEAMBIMCIVIL S.L.", "TEAMBIMCIVIL S.L. · Director de Obra"]
                    and en_lista == "TEAMBIMCIVIL S.L. · Director de Obra", f"{datos} {en_lista}")
    caso(S2, s2)

    # ── S3 · Una revisión importada que corrige el expediente: el mal escrito deja de proponerse ──
    S3 = "Una revisión importada que corrige el expediente: el mal escrito deja de proponerse, también tras borrar la Rev. 0"

    def s3():
        ctx, pg = telefono([{"id": "s3-r0", "num": "425_2026.COM-AV-01", "expediente": "425_2026.COM", "momento": "2026-10-05T09:00:00.000Z"}],
                           recordados={"expedientes": [{"valor": "425_2026.COM", "fecha": "2026-10-05T09:00:00.000Z", "actas": ["s3-r0"]}],
                                       "personas": []})
        m = "2026-10-06T09:00:00.000Z"
        revision = {"id": "s3-r1", "esquema": 4, "estado": "emitida", "creada": m, "modificada": m, "emitida": m, "versionEmision": "3.1.3",
                    "revision": 1, "fase": 8, "corrigeA": "s3-r0",
                    "revisiones": [{"rev": 0, "fecha": "2026-10-05T09:00:00.000Z", "motivo": "Emisión inicial"},
                                   {"rev": 1, "fecha": m, "motivo": "Expediente mal escrito"}],
                    "datos": {"actaNum": "425_2026.COM-AV-01", "numAuto": False, "tecnico": "Blas González", "cargoTecnico": "Director de Obra",
                              "expediente": "425_2026.CON", "objeto": "tajos", "descripcion": "", "fecha": "2026-10-05", "hora": "09:00"},
                    "gps": {"lat": "37.389210", "lng": "-5.984460", "acc": 8, "hora": m},
                    "sinGps": {"activo": False, "referencia": "", "lat": "", "lng": ""}, "condiciones": {"meteo": [], "temperatura": ""},
                    "asistentes": [], "observaciones": "", "seguimiento": [], "seguimientoOrigen": None, "seguimientoNotas": [],
                    "seguimientoGuardado": None, "indicaciones": [], "fotos": [], "firmasAsistentes": {},
                    "firma": [{"refW": 300, "refH": 150, "puntos": [[10, 10], [60, 40], [120, 20], [180, 60]]}]}
        importar(pg, zip_de_acta(SALIDA / "S313-revision-importada.zip", revision))
        pg.wait_for_function("Array.from(document.getElementById('expedienteLista').options).some(o => o.value === '425_2026.CON')",
                             timeout=5000)
        con_rev0 = opciones(pg, "expedienteLista")
        pg.evaluate("Almacen.borrarActa('s3-r0')")
        pg.evaluate("cerrarMisActas()")
        pg.wait_for_timeout(500)
        sin_rev0 = opciones(pg, "expedienteLista")
        r.comprobar(S3, con_rev0 == sin_rev0 == ["", "425_2026.CON", "__nuevo"], f"{con_rev0} {sin_rev0}")
    caso(S3, s3)

    # ── S4 · La empresa escrita en una ficha se ofrece al momento en las demás; «Otra…» parte siempre del valor actual ──
    S4 = "Empresa: la escrita en una ficha se ofrece al momento en las demás que ya había, y «Otra…» parte siempre del valor actual"

    def s4():
        ctx, pg = telefono(una)
        pg.select_option("#expedienteLista", "EXP-BASE")
        pg.select_option("#tecnico", "Blas González")
        pg.evaluate("irAFase(2)")
        anadir_asistente(pg)
        pg.fill("#asNombre0", "Juan Ruiz")
        anadir_asistente(pg)
        pg.fill("#asNombre1", "Luis Gil")
        pg.select_option("#asEmpresaLista0", "__otro")
        pg.fill("#asEmpresa0", "Constructora Norte")
        pg.click("#asNombre1")                                # sale del campo
        en_otra = "Constructora Norte" in opciones(pg, "asEmpresaLista1")
        pg.select_option("#asEmpresaLista1", "__otro")
        pg.fill("#asEmpresa1", "Vieja S.A.")
        pg.select_option("#asEmpresaLista1", "TEAMBIMCIVIL S.L.")
        pg.select_option("#asEmpresaLista1", "__otro")
        actual = pg.evaluate("[document.getElementById('asEmpresa1').value, state.acta.asistentes[1].empresa]")
        r.comprobar(S4, en_otra and actual == ["TEAMBIMCIVIL S.L.", "TEAMBIMCIVIL S.L."], f"{en_otra} {actual}")
    caso(S4, s4)

    # ── S5 · Un acta guardada con una versión más nueva de la app (otra ventana actualizada) no se abre con esta ──
    S5 = "Un acta guardada con una versión más nueva de la app no se abre con esta: avisa para actualizar y no se queda en blanco"

    def s5():
        ctx, pg = telefono()
        esquema, ultima = pg.evaluate("[ESQUEMA, ULTIMA_FASE]")              # una versión posterior: más fases y otro esquema
        pg.evaluate(INYECTAR, [{"id": "s5-nueva", "num": "EXP-S5-AV-01", "expediente": "EXP-S5", "momento": "2026-10-07T09:00:00.000Z",
                                "estado": "borrador", "fase": ultima + 1, "esquema": esquema + 1}])
        pg.evaluate("localStorage.setItem('tbc-acta-actual', 's5-nueva')")
        pg.reload()
        pg.wait_for_function("state.acta !== null && document.getElementById('alertOverlay').classList.contains('open')", timeout=10000)
        aviso1 = cerrar_aviso(pg) or ""
        al_abrir = pg.evaluate("""[state.acta.id !== 's5-nueva', state.persistida, state.fase,
            !document.getElementById('avisoVersion').classList.contains('hidden'),
            Array.from(document.querySelectorAll('.phase-container.active')).map(s => s.id)]""")
        pg.evaluate("abrirActaGuardada('s5-nueva')")
        aviso2 = cerrar_aviso(pg) or ""
        r.comprobar(S5, "versión más nueva" in aviso1 and al_abrir == [True, False, 1, True, ["phase1"]]
                    and "versión más nueva" in aviso2 and pg.evaluate("state.acta.id") != "s5-nueva", f"{al_abrir} | {aviso2[:40]}")
    caso(S5, s5)

    # ── PDF de las actas emitidas con la v3.1.2 y con la v3.1.3: idéntico al que se distribuyó ──
    for version, datos in (("3.1.2", DATOS_V312), ("3.1.3", DATOS_V313)):
        nombre = f"Actas emitidas con la v{version}: su PDF es idéntico al que se distribuyó (texto y contenido de cada página)"

        def pdf_iguales(datos=datos, version=version, nombre=nombre):
            ctx, pg = telefono()
            esperadas = json.loads((datos / "huellas-pdf.json").read_text(encoding="utf-8"))
            importar(pg, datos / "av01.zip")
            importar(pg, datos / "av02.zip")
            iguales = {}
            for n in ("av01", "av02"):
                with zipfile.ZipFile(datos / f"{n}.zip") as z:
                    numero = json.loads(z.read(next(x for x in z.namelist() if x.endswith("acta.json"))))["acta"]["datos"]["actaNum"]
                acta_id = pg.evaluate(f"Almacen.listarActas().then(l => l.find(a => a.datos.actaNum === '{numero}').id)")
                pg.evaluate(f"abrirActaGuardada('{acta_id}')")
                pg.wait_for_function(f"state.acta.id === '{acta_id}' && state.fase === ULTIMA_FASE")
                pg.wait_for_function("document.getElementById('emitPdfEstado').textContent.startsWith('PDF listo')", timeout=60000)
                ruta = SALIDA / f"S313-v{version.replace('.', '')}-{n}.pdf"
                descargar(pg, "#btnDescargar", ruta)
                esperado = (datos / f"pdf-{n}.txt").read_text(encoding="utf-8").strip()
                iguales[n] = [texto_pdf(ruta) == esperado, huellas(ruta) == esperadas[n]]
            r.comprobar(nombre, all(all(v) for v in iguales.values()), str(iguales))
        caso(nombre, pdf_iguales)


def main():
    r = Resultados()
    SALIDA.mkdir(parents=True, exist_ok=True)
    horizontal, _, _ = (str(f) for f in generar(SALIDA / "fotos"))
    srv = servidor()
    externas, errores = [], []

    with sync_playwright() as p:
        navegador = p.chromium.launch()
        ctx = contexto_movil(p, navegador)
        vigilar_red(ctx, externas)
        pg = abrir(ctx, errores)
        cdp = ctx.new_cdp_session(pg)
        esquema = pg.evaluate("ESQUEMA")

        # ══ Mejora 1 · Cabecera en gris corporativo ══
        estilos = pg.evaluate(ESTILOS)
        manifiesto = json.loads((RAIZ / "manifest.webmanifest").read_text(encoding="utf-8"))
        r.comprobar("Cabecera y pasos en gris corporativo #888888; también la barra del navegador y la app instalada",
                    estilos["cabecera"] == GRIS and estilos["pasos"] == GRIS and estilos["tema"].lower() == "#888888"
                    and manifiesto["theme_color"].lower() == "#888888", str({k: estilos[k] for k in ("cabecera", "pasos", "tema")}))
        contrastes = {k: round(contraste(estilos[k], GRIS), 2) for k in ("titulo", "subtitulo", "guardado", "paso", "activo")}
        r.comprobar("Sobre el gris, todos los textos en negro (≥ 4,5:1), también el título y sin sombra",
                    all(v >= 4.5 for v in contrastes.values()) and rgb(estilos["titulo"]) == [26, 26, 26]
                    and estilos["sombra"] == "none", str(contrastes))
        version = pg.evaluate("VERSION")
        r.comprobar("La app se llama «Acta de Campo»: cabecera, pestaña del navegador y app instalada",
                    estilos["nombre"] == f"Acta de Campo · v{version}" and estilos["pestana"] == "Acta de Campo · TeamBIMCivil"
                    and manifiesto["name"] == "Acta de Campo · TEAMBIMCIVIL" and "visita" not in manifiesto["description"].lower(),
                    f"{estilos['nombre']} | {estilos['pestana']} | {manifiesto['name']}")
        abrir_mis_actas(pg)
        panel = pg.evaluate("""[getComputedStyle(document.querySelector('.panel-header')).backgroundColor,
                                getComputedStyle(document.querySelector('.panel-title')).color]""")
        r.comprobar("Mis actas: cabecera del panel en el mismo gris, con el título legible",
                    panel[0] == GRIS and contraste(panel[1], GRIS) >= 4.5, str(panel))
        pg.screenshot(path=str(SALIDA / "s31-mis-actas.png"))
        pg.evaluate("cerrarMisActas()")

        # ══ Mejora 2 · El expediente, primer campo ══
        orden = pg.evaluate("Array.from(document.querySelectorAll('#phase1 input, #phase1 select')).map(e => e.id).filter(Boolean)")
        esperado = ["expedienteLista", "expediente", "actaNum", "actaRev", "tecnico", "cargoTecnico", "cargoTecnicoOtro", "objeto",
                    "descripcion", "fechaVisita", "horaVisita"]
        r.comprobar("Identificación: expediente, nº de acta, técnico, cargo, objeto y detalle, fecha y hora, en ese orden",
                    orden[:len(esperado)] == esperado, ", ".join(orden[:11]))

        # ══ Mejoras 3, 4 y 5 · Listas desplegables ══
        tecnicos, cargos, objetos = (pg.evaluate(OPCIONES, i) for i in ("tecnico", "cargoTecnico", "objeto"))
        r.comprobar("Técnicos de TEAMBIMCIVIL en una lista, en orden alfabético y con tildes",
                    tecnicos == [["", "Elige el técnico"]] + [[t, t] for t in TECNICOS], ", ".join(t[1] for t in tecnicos))
        r.comprobar("Cargo del técnico: los 4 roles, en su orden, y «Otro»",
                    cargos == [["", "Elige el cargo"]] + [[c, c] for c in CARGOS] + [["__otro", "Otro…"]], ", ".join(c[1] for c in cargos))
        r.comprobar("Objeto de la visita: los 5 de la lista",
                    [o[1] for o in objetos] == ["Elige el objeto"] + OBJETOS and objetos[-1][0] == "otro", ", ".join(o[1] for o in objetos))

        def aviso_al_seguir():
            pg.click("#btnNext")
            return cerrar_aviso(pg) or ""

        avisos = [aviso_al_seguir()]                          # nada rellenado
        escribir_expediente(pg, EXP)
        pg.wait_for_function(f"state.acta.datos.actaNum === '{EXP}-AV-01'", timeout=5000)
        avisos.append(aviso_al_seguir())                      # falta el técnico
        pg.select_option("#tecnico", "Ana Cascajosa")
        avisos.append(aviso_al_seguir())                      # falta el cargo
        pg.select_option("#cargoTecnico", "__otro")
        otro_cargo = [pg.is_visible("#cargoTecnicoOtro"), pg.evaluate("document.activeElement.id")]
        avisos.append(aviso_al_seguir())                      # «Otro» sin escribir
        pg.fill("#cargoTecnicoOtro", "Jefe de equipo")
        avisos.append(aviso_al_seguir())                      # falta el objeto
        pg.select_option("#objeto", "otro")
        otro_objeto = [pg.text_content("#descripcionEtiqueta"), pg.evaluate("document.activeElement.id")]
        avisos.append(aviso_al_seguir())                      # «Otro» sin describir
        pg.fill("#descripcion", DAÑOS)
        avisos.append(aviso_al_seguir())                      # falta la ubicación
        claves = ["expediente", "técnico", "cargo", "cargo", "objeto", "describe el objeto", "ubicación"]
        r.comprobar("Avisos en el orden de los campos: expediente, técnico, cargo (también «Otro» vacío), objeto, «Otro» sin describir, ubicación",
                    len(avisos) == len(claves) and all(c in a.lower() for c, a in zip(claves, avisos)), " | ".join(a.split("|")[-1][:34] for a in avisos))
        r.comprobar("Cargo «Otro»: aparece un campo para escribirlo (con el cursor en él)",
                    otro_cargo == [True, "cargoTecnicoOtro"] and pg.evaluate("state.acta.datos.cargoTecnico") == "Jefe de equipo", str(otro_cargo))
        r.comprobar("Objeto «Otro»: el detalle pasa a «Describe el objeto», obligatorio",
                    "Describe el objeto" in otro_objeto[0] and otro_objeto[1] == "descripcion", str(otro_objeto))
        pg.click("#gpsBtn")
        pg.wait_for_selector("#gpsStatus.ok", timeout=15000)
        pg.click("#btnNext")                                  # → asistentes

        # ── Asistentes: el técnico con su cargo y el cargo de cada asistente de la lista ──
        r.comprobar("Asistentes: el técnico con su cargo, que ya no se pide en esta fase",
                    pg.evaluate("state.fase") == 2 and pg.text_content("#tecnicoNombreAsist") == "Ana Cascajosa"
                    and pg.text_content("#tecnicoCargoAsist") == "TEAMBIMCIVIL S.L. · Jefe de equipo"
                    and pg.locator("#phase2 select:not(#anadirAsistenteLista), #phase2 input").count() == 0)
        anadir_asistente(pg)
        pg.fill("#asNombre0", "Juan Pérez")
        escribir_empresa(pg, 0, "Constructora Sur")
        lista_asistente = pg.evaluate(OPCIONES, "asCargo0")
        pg.select_option("#asCargo0", "Supervisor de Obra")
        de_lista = pg.evaluate("[state.acta.asistentes[0].cargo, document.getElementById('asCargoOtro0').classList.contains('hidden')]")
        pg.select_option("#asCargo0", "__otro")
        foco = pg.evaluate("document.activeElement.id")
        pg.fill("#asCargoOtro0", "Encargado")
        anadir_asistente(pg)       # se vuelve a pintar la lista
        pg.fill("#asNombre1", "Marta Ruiz")
        escribir_empresa(pg, 1, "Propiedad")
        pg.select_option("#asCargo1", "Redactor de Estudio")
        repintado = pg.evaluate("""[document.getElementById('asCargo0').value, document.getElementById('asCargoOtro0').value,
            document.getElementById('asCargoOtro0').classList.contains('hidden'), state.acta.asistentes.map(a => a.cargo)]""")
        r.comprobar("Cargo de los asistentes: la misma lista; uno de ella o «Otro» escrito, que se conserva al repintar",
                    lista_asistente == cargos and de_lista == ["Supervisor de Obra", True] and foco == "asCargoOtro0"
                    and repintado == ["__otro", "Encargado", False, ["Encargado", "Redactor de Estudio"]], str(repintado))
        pg.screenshot(path=str(SALIDA / "s31-asistentes.png"), full_page=True)
        pg.click("#btnNext")                                  # → fotos
        pg.set_input_files("#photoInput", [horizontal])
        esperar_fotos(pg, 1)
        pg.click("#btnNext")                                  # → notas
        pg.click("#btnNext")                                  # → acuerdos
        responsables = pg.evaluate("responsablesPosibles()")
        r.comprobar("Responsables posibles con los cargos de la lista y los escritos",
                    responsables == ["TEAMBIMCIVIL S.L. · Jefe de equipo", "Constructora Sur · Encargado", "Propiedad · Redactor de Estudio"],
                    str(responsables))
        pg.click("#btnNext")                                  # → firmas
        firmar(cdp, pg)
        pg.once("dialog", lambda dialogo: dialogo.accept())   # los asistentes no firman: «No firma»
        pg.click("#btnNext")                                  # → revisión
        general = pg.inner_text("#reviewGeneral")
        r.comprobar("Revisión: con «Otro», el objeto es lo descrito", f"Objeto\t{DAÑOS}" in general or f"Objeto\n{DAÑOS}" in general,
                    general.replace("\n", " | ")[:160])
        pg.click("#btnNext")                                  # → emisión
        mensaje = pg.evaluate("textoCompartir()")
        ruta1 = SALIDA / "S31-AV01.pdf"
        descargar(pg, "#btnDescargar", ruta1)
        pg.wait_for_function("state.acta.estado === 'emitida'")
        acta1 = pg.evaluate("state.acta.id")
        texto1 = texto_pdf(ruta1)
        recordados = pg.evaluate("JSON.parse(localStorage.getItem('tbc-recordados') || 'null')") or {}
        r.comprobar("Al emitir, el teléfono recuerda el expediente y los asistentes (para proponerlos aunque se borre el acta)",
                    [e["valor"] for e in recordados.get("expedientes", [])] == [EXP]
                    and [(x["nombre"], x["empresa"], x["cargo"], x["expedientes"]) for x in recordados.get("personas", [])]
                    == [("Juan Pérez", "Constructora Sur", "Encargado", [EXP]), ("Marta Ruiz", "Propiedad", "Redactor de Estudio", [EXP])],
                    str(recordados)[:200])
        r.comprobar("PDF: el técnico de «TEAMBIMCIVIL S.L.» en Asistentes, como en Firmas y en las propiedades",
                    "Ana Cascajosa TEAMBIMCIVIL S.L. Jefe de equipo" in texto1
                    and (PdfReader(str(ruta1)).metadata or {}).get("/Author") == "Ana Cascajosa · TEAMBIMCIVIL S.L.",
                    texto1[texto1.find("ASISTENTES"):texto1.find("ASISTENTES") + 120])
        r.comprobar("PDF: técnico, objeto descrito y cargos de la lista o escritos",
                    f"OBJETO {DAÑOS}" in texto1 and "Otro ·" not in texto1 and "TÉCNICO Ana Cascajosa" in texto1
                    and all(s in texto1 for s in ("Jefe de equipo", "Encargado", "Redactor de Estudio")) and f"Objeto: {DAÑOS}" in mensaje)
        lector1 = PdfReader(str(ruta1))
        paginas1, meta1 = len(lector1.pages), lector1.metadata or {}
        r.comprobar("Formato del SGC «Acta de Campo»: título, bloque de control de cada página, propiedades y texto al compartir",
                    "ACTA DE CAMPO" in texto1 and texto1.count("Acta de Campo") == paginas1 and "de Inspección" not in texto1
                    and "DE INSPECCIÓN" not in texto1 and meta1.get("/Title") == f"Acta de Campo {EXP}-AV-01 · Rev. 0"
                    and "acta de campo" in (meta1.get("/Keywords") or "") and (meta1.get("/Creator") or "").startswith("Acta de Campo v")
                    and mensaje.startswith(f"Acta de Campo {EXP}-AV-01 · Rev. 0"), f"{meta1.get('/Title')} | {mensaje[:40]}")
        r.comprobar("PDF: cabecera y pie en gris #888888 en todas las páginas",
                    paginas1 >= 2 and marco_pdf(ruta1) == [MARCO_GRIS] * paginas1, str(marco_pdf(ruta1)))
        r.comprobar("Un acta emitida: técnico, cargo y objeto bloqueados",
                    pg.evaluate("['tecnico', 'cargoTecnico', 'cargoTecnicoOtro', 'objeto', 'descripcion'].every(id => document.getElementById(id).disabled)"))

        # ══ Propuestas: técnico y cargo recordados; sugerencias de expedientes ══
        abrir_mis_actas(pg)
        pg.click("#panelActas button:has-text('Nueva acta')")
        pg.wait_for_function("state.persistida === false && state.acta.datos.expediente === ''")
        nueva = pg.evaluate("""['tecnico', 'cargoTecnico', 'cargoTecnicoOtro', 'objeto', 'expediente']
            .map(id => document.getElementById(id).value)""")
        r.comprobar("Un acta nueva propone el último técnico y cargo del teléfono (no el objeto ni el expediente)",
                    nueva == ["Ana Cascajosa", "__otro", "Jefe de equipo", "", ""], str(nueva))
        pg.wait_for_function(f"""!document.getElementById('expedienteLista').classList.contains('hidden')
            && Array.from(document.getElementById('expedienteLista').options).some(o => o.value === '{EXP}')""", timeout=5000)
        r.comprobar("El expediente se elige de un desplegable con los ya usados en el teléfono (sin listas del navegador)",
                    pg.evaluate(OPCIONES, "expedienteLista") == [["", "Elige el expediente"], [EXP, EXP], ["__nuevo", "➕ Nuevo expediente…"]]
                    and pg.is_hidden("#expediente") and pg.locator("datalist").count() == 0)

        # ── «Nueva visita»: copia técnico, cargo, objeto y detalle; objeto de la lista con detalle ──
        abrir_mis_actas(pg)
        pg.click(f".acta-card[data-acta='{acta1}'] button:has-text('Nueva visita')")
        pg.wait_for_function(f"state.acta.datos.actaNum === '{EXP}-AV-02'", timeout=10000)
        copia = pg.evaluate("[state.acta.datos.tecnico, state.acta.datos.cargoTecnico, state.acta.datos.objeto, state.acta.datos.descripcion]")
        r.comprobar("«Nueva visita» copia el técnico, su cargo y el objeto con su detalle",
                    copia == ["Ana Cascajosa", "Jefe de equipo", "otro", DAÑOS], str(copia))
        pg.select_option("#tecnico", "Blas González")
        pg.select_option("#cargoTecnico", "Director de Obra")
        pg.select_option("#objeto", "tajos")
        pg.fill("#descripcion", LOSA)
        r.comprobar("Objeto de la lista con detalle: «Inspección de Tajos de Obra · detalle»",
                    pg.evaluate("textoObjeto(state.acta)") == f"Inspección de Tajos de Obra · {LOSA}"
                    and pg.is_hidden("#cargoTecnicoOtro") and "Detalle (opcional)" in pg.text_content("#descripcionEtiqueta"))
        pg.click("#gpsBtn")
        pg.wait_for_selector("#gpsStatus.ok", timeout=15000)
        pg.evaluate("window.scrollTo(0, 0)")
        pg.screenshot(path=str(SALIDA / "s31-identificacion.png"))
        pg.evaluate("guardarYa()")

        # ── Registro en Excel: objeto y detalle en columnas propias, técnico y cargo ──
        abrir_mis_actas(pg)
        descargar(pg, "#btnRegistro", SALIDA / "S31-registro.xlsx")
        cerrar_aviso(pg)
        hoja = load_workbook(SALIDA / "S31-registro.xlsx")["Actas"]
        cabecera = [c.value for c in hoja[1]]
        filas = [f[6:10] for f in hoja.iter_rows(min_row=2, values_only=True)]
        r.comprobar("Registro: columnas Objeto, Detalle, Técnico y Cargo del técnico con los valores de cada acta",
                    cabecera[6:10] == ["Objeto", "Detalle", "Técnico", "Cargo del técnico"]
                    and filas == [("Otro", DAÑOS, "Ana Cascajosa", "Jefe de equipo"),
                                  ("Inspección de Tajos de Obra", LOSA, "Blas González", "Director de Obra")], str(filas))
        pg.evaluate("cerrarMisActas()")
        ctx.close()

        # ══ Actas de la v3.0 (esquema 3): migración y PDF sin cambios ══
        ctx_c = contexto_movil(p, navegador)
        vigilar_red(ctx_c, externas)
        pc = abrir(ctx_c, errores)
        pc.evaluate("""async () => {
            const base = { esquema: 3, creada: '2026-10-20T08:00:00.000Z', modificada: '2026-10-20T09:00:00.000Z', revision: 0,
              datos: { actaNum: 'V30-AV-01', numAuto: false, tecnico: 'Blas González González', cargoTecnico: 'Dirección facultativa',
                       expediente: 'V30', descripcion: 'Texto libre de la v3.0', fecha: '2026-10-20', hora: '09:00' },
              gps: { lat: '37.389210', lng: '-5.984460', acc: 8, hora: '2026-10-20T08:05:00.000Z' },
              sinGps: { activo: false, referencia: '', lat: '', lng: '' }, condiciones: { meteo: [], temperatura: '' },
              asistentes: [], observaciones: '', seguimiento: [], seguimientoOrigen: null, seguimientoNotas: [],
              seguimientoGuardado: null, indicaciones: [], fotos: [], firmasAsistentes: {},
              firma: [{ refW: 300, refH: 150, puntos: [[10, 10], [60, 40], [120, 20], [180, 60], [240, 30]] }] };
            await Almacen.guardarActa({ ...base, id: 'acta-v30-emitida', estado: 'emitida', emitida: '2026-10-20T10:00:00.000Z',
              versionEmision: '3.0', fase: 8, revisiones: [{ rev: 0, fecha: '2026-10-20T10:00:00.000Z', motivo: 'Emisión inicial' }] });
            await Almacen.guardarActa({ ...base, id: 'acta-v30-borrador', estado: 'borrador', emitida: null, fase: 1,
              revisiones: [{ rev: 0, fecha: null, motivo: 'Emisión inicial' }], datos: { ...base.datos, actaNum: 'V30-AV-02' } });
            await Almacen.guardarActa({ ...base, esquema: 4, id: 'acta-v31-emitida', estado: 'emitida', emitida: '2026-10-23T16:00:00.000Z',
              versionEmision: '3.1', fase: 8, revisiones: [{ rev: 0, fecha: '2026-10-23T16:00:00.000Z', motivo: 'Emisión inicial' }],
              datos: { ...base.datos, actaNum: 'V31-AV-01', expediente: 'V31', tecnico: 'Blas González', cargoTecnico: 'Director de Obra',
                       objeto: 'tajos', descripcion: 'Losa de planta 2', fecha: '2026-10-23', hora: '17:30' } });
            localStorage.setItem('tbc-acta-actual', 'acta-v30-borrador');
        }""")
        pc.reload()
        pc.wait_for_function("state.acta && state.acta.id === 'acta-v30-borrador'")
        migrada = pc.evaluate("""({ esquema: state.acta.esquema, objeto: state.acta.datos.objeto,
            tecnico: document.getElementById('tecnico').value,
            opcion: Array.from(document.getElementById('tecnico').options).some(o => o.value === 'Blas González González'),
            cargo: [document.getElementById('cargoTecnico').value, document.getElementById('cargoTecnicoOtro').value],
            detalle: document.getElementById('descripcion').value })""")
        r.comprobar("Un borrador de la v3.0 se migra: su técnico y su cargo, que no están en las listas, se muestran tal cual",
                    migrada == {"esquema": esquema, "objeto": "", "tecnico": "Blas González González", "opcion": True,
                                "cargo": ["__otro", "Dirección facultativa"], "detalle": "Texto libre de la v3.0"}, str(migrada))
        pc.click("#btnNext")
        aviso = cerrar_aviso(pc) or ""
        r.comprobar("…y al seguir pide elegir el objeto de la visita (el texto anterior queda como detalle)",
                    "Elige el objeto" in aviso and pc.evaluate("state.fase") == 1, aviso)
        pc.evaluate("abrirActaGuardada('acta-v30-emitida')")
        pc.wait_for_function("state.acta.id === 'acta-v30-emitida' && state.fase === 8")
        pc.wait_for_function("document.getElementById('emitPdfEstado').textContent.startsWith('PDF listo')", timeout=60000)
        descargar(pc, "#btnDescargar", SALIDA / "S31-v30-emitida.pdf")
        texto_v30 = texto_pdf(SALIDA / "S31-v30-emitida.pdf")
        r.comprobar("Un acta emitida con la v3.0 genera su PDF como entonces (objeto con su texto, técnico, empresa y cargo)",
                    "OBJETO Texto libre de la v3.0" in texto_v30 and "TÉCNICO Blas González González" in texto_v30
                    and "Blas González González TEAMBIMCIVIL Dirección facultativa" in texto_v30 and "S.L." not in texto_v30)
        lector_v30 = PdfReader(str(SALIDA / "S31-v30-emitida.pdf"))
        paginas_v30, marco_v30 = len(lector_v30.pages), marco_pdf(SALIDA / "S31-v30-emitida.pdf")
        r.comprobar("…y con su formato: «Acta de Visita de Inspección», cabecera y pie en negro, también al compartir",
                    "ACTA DE VISITA DE INSPECCIÓN" in texto_v30 and texto_v30.count("Acta de Visita de Inspección") == paginas_v30
                    and "Acta de Campo" not in texto_v30 and "ACTA DE CAMPO" not in texto_v30 and marco_v30 == [MARCO_NEGRO] * paginas_v30
                    and (lector_v30.metadata or {}).get("/Title") == "Acta de Visita de Inspección V30-AV-01 · Rev. 0"
                    and pc.evaluate("textoCompartir()").startswith("Acta de Visita de Inspección V30-AV-01 · Rev. 0"), str(marco_v30))
        # La v3.1 (publicada el 23/10/2026) emitía con el formato anterior: sus actas lo conservan en la v3.1.1.
        pc.evaluate("abrirActaGuardada('acta-v31-emitida')")
        pc.wait_for_function("state.acta.id === 'acta-v31-emitida' && state.fase === 8")
        pc.wait_for_function("document.getElementById('emitPdfEstado').textContent.startsWith('PDF listo')", timeout=60000)
        descargar(pc, "#btnDescargar", SALIDA / "S31-v31-emitida.pdf")
        texto_v31, marco_v31 = texto_pdf(SALIDA / "S31-v31-emitida.pdf"), marco_pdf(SALIDA / "S31-v31-emitida.pdf")
        r.comprobar("Un acta emitida con la v3.1 conserva su PDF: «Acta de Visita de Inspección», marco negro y enlaces al mapa",
                    "ACTA DE VISITA DE INSPECCIÓN" in texto_v31 and "Acta de Campo" not in texto_v31 and "ACTA DE CAMPO" not in texto_v31
                    and marco_v31 == [MARCO_NEGRO] * len(marco_v31) and "Ver en el mapa" in texto_v31
                    and "OBJETO Inspección de Tajos de Obra · Losa de planta 2" in texto_v31
                    and pc.evaluate("textoCompartir()").startswith("Acta de Visita de Inspección V31-AV-01 · Rev. 0"), str(marco_v31))
        importado = pc.evaluate("""(() => {
            const ficha = { formato: 'tbc-acta', acta: { id: 'acta-imp-1', estado: 'borrador', datos: { objeto: 'patologia' }, fotos: [] } };
            const valido = actaDesdeArchivo(ficha).datos.objeto;
            ficha.acta.datos.objeto = '<b>xyz</b>';
            const desconocido = actaDesdeArchivo(ficha).datos.objeto;
            ficha.acta.datos.objeto = 5;
            return [valido, desconocido, actaDesdeArchivo(ficha).datos.objeto];
        })()""")
        r.comprobar("Importación: solo se acepta un objeto de la lista", importado == ["patologia", "", ""], str(importado))
        ctx_c.close()

        # ══ v3.1.3 · Desplegables con lo usado en el teléfono: expediente, asistentes, su empresa y el responsable ══
        ctx_d = contexto_movil(p, navegador)
        vigilar_red(ctx_d, externas)
        pd = abrir(ctx_d, errores)
        pd.evaluate("""async () => {
            const asis = (id, nombre, empresa, cargo, firma = true) => ({ id, nombre, empresa, cargo, firma });
            const acta = (id, num, expediente, momento, asistentes) => ({
              id, esquema: 4, estado: 'emitida', creada: momento, modificada: momento, emitida: momento, versionEmision: '3.1.2',
              revision: 0, fase: 8, revisiones: [{ rev: 0, fecha: momento, motivo: 'Emisión inicial' }],
              datos: { actaNum: num, numAuto: false, tecnico: 'Blas González', cargoTecnico: 'Director de Obra', expediente,
                       objeto: 'tajos', descripcion: '', fecha: momento.slice(0, 10), hora: '09:00' },
              gps: { lat: '37.389210', lng: '-5.984460', acc: 8, hora: momento }, sinGps: { activo: false, referencia: '', lat: '', lng: '' },
              condiciones: { meteo: [], temperatura: '' }, asistentes, observaciones: '', seguimiento: [], seguimientoOrigen: null,
              seguimientoNotas: [], seguimientoGuardado: null, indicaciones: [], fotos: [], firmasAsistentes: {},
              firma: [{ refW: 300, refH: 150, puntos: [[10, 10], [60, 40], [120, 20], [180, 60]] }] });
            await Almacen.guardarActa(acta('d-a1', 'EXP-A-AV-01', 'EXP-A', '2026-10-01T09:00:00.000Z',
              [asis('x1', 'Juan Pérez', 'Constructora Sur', 'Jefe de obra'), asis('x2', 'Ana Cascajosa', 'teambimcivil', 'Redactor de Estudio')]));
            await Almacen.guardarActa(acta('d-a2', 'EXP-A-AV-02', ' Exp-A ', '2026-10-08T09:00:00.000Z',
              [asis('x3', 'Juan Pérez', 'Constructora Sur', 'Encargado', false)]));
            await Almacen.guardarActa(acta('d-b1', 'EXP-B-AV-01', 'EXP-B', '2026-10-15T09:00:00.000Z',
              [asis('x4', 'Marta Ruiz', 'Propiedad', 'Promotora'), asis('x5', 'Luis Gómez', 'CSS Ingeniería', '')]));
            localStorage.setItem('tbc-recordados', JSON.stringify({
              expedientes: [{ valor: 'EXP-RECORDADO', fecha: '2026-09-01T09:00:00.000Z' }],
              personas: [{ nombre: 'Rosa Gil', empresa: 'Geotecnia Andaluza', cargo: 'Geóloga', firma: true,
                           fecha: '2026-09-01T09:00:00.000Z', expedientes: ['EXP-RECORDADO'] }] }));
        }""")
        pd.reload()
        pd.wait_for_function("state.acta && !state.persistida && !document.getElementById('expedienteLista').classList.contains('hidden')",
                             timeout=10000)
        lista_exp = pd.evaluate(OPCIONES, "expedienteLista")
        r.comprobar("Expediente: los usados en el teléfono y los recordados, del más reciente al más antiguo, sin repetir y con «Nuevo»",
                    lista_exp == [["", "Elige el expediente"], ["EXP-B", "EXP-B"], ["Exp-A", "Exp-A"],
                                  ["EXP-RECORDADO", "EXP-RECORDADO"], ["__nuevo", "➕ Nuevo expediente…"]]
                    and pd.is_hidden("#expediente") and pd.evaluate("document.getElementById('expedienteEtiqueta').htmlFor") == "expedienteLista",
                    str(lista_exp))
        pd.select_option("#expedienteLista", "Exp-A")
        pd.wait_for_function("state.acta.datos.actaNum.toLowerCase() === 'exp-a-av-03'", timeout=5000)
        elegido = pd.evaluate("[state.acta.datos.expediente, state.acta.datos.actaNum]")
        pd.select_option("#expedienteLista", "__nuevo")
        pd.wait_for_function("state.acta.datos.actaNum === ''", timeout=5000)
        nuevo = pd.evaluate("""[document.activeElement.id, state.acta.datos.expediente, state.acta.datos.actaNum,
            document.getElementById('expedienteEtiqueta').htmlFor]""") + [pd.is_visible("#expediente")]
        pd.fill("#expediente", "EXP-C")
        pd.wait_for_function("state.acta.datos.actaNum === 'EXP-C-AV-01'", timeout=5000)
        pd.select_option("#expedienteLista", "EXP-B")
        pd.wait_for_function("state.acta.datos.actaNum === 'EXP-B-AV-02'", timeout=5000)
        r.comprobar("Elegir un expediente propone su número; «Nuevo expediente…» abre el campo para escribirlo, y se puede volver a la lista",
                    elegido == ["Exp-A", "Exp-A-AV-03"] and nuevo == ["expediente", "", "", "expediente", True]
                    and pd.is_hidden("#expediente") and pd.evaluate("state.acta.datos.expediente") == "EXP-B", f"{elegido} {nuevo}")

        pd.select_option("#tecnico", "Blas González")
        pd.select_option("#cargoTecnico", "Director de Obra")
        pd.select_option("#objeto", "tajos")
        pd.click("#gpsBtn")
        pd.wait_for_selector("#gpsStatus.ok", timeout=15000)
        pd.click("#btnNext")                                  # → asistentes
        pd.wait_for_function("state.fase === 2 && document.querySelectorAll('#anadirAsistenteLista optgroup').length === 3")
        menu = pd.evaluate(GRUPOS)
        tecnicos_menu = [t for t in TECNICOS if t != "Blas González"]
        esperado_menu = {
            "sueltas": ["➕ Añadir asistente…", "✏️ Persona nueva (escribir)"],
            "grupos": [["TEAMBIMCIVIL S.L.", [f"{t} · TEAMBIMCIVIL S.L." + (" · Redactor de Estudio" if t == "Ana Cascajosa" else "")
                                              for t in tecnicos_menu]],
                       ["De este expediente", ["Marta Ruiz · Propiedad · Promotora", "Luis Gómez · CSS Ingeniería"]],
                       ["De otras actas", ["Juan Pérez · Constructora Sur · Encargado", "Rosa Gil · Geotecnia Andaluza · Geóloga"]]]}
        r.comprobar("«➕ Añadir asistente…»: los técnicos de TEAMBIMCIVIL S.L. (menos el del acta) y las personas de actas anteriores, primero las del expediente",
                    menu == esperado_menu, str(menu)[:300])
        anadir_asistente(pd, "Ana Cascajosa · TEAMBIMCIVIL S.L. · Redactor de Estudio")
        anadir_asistente(pd, "Juan Pérez · Constructora Sur · Encargado")
        fichas = pd.evaluate("""state.acta.asistentes.map((a, i) => [a.nombre, a.empresa, a.cargo, a.firma,
            document.getElementById('asEmpresaLista' + i).value, document.getElementById('asEmpresa' + i).classList.contains('hidden')])""")
        menu2 = pd.evaluate(GRUPOS)
        r.comprobar("Elegido de la lista: la ficha sale con su empresa (de su desplegable), cargo y firma, y deja de proponerse",
                    fichas == [["Ana Cascajosa", "TEAMBIMCIVIL S.L.", "Redactor de Estudio", True, "TEAMBIMCIVIL S.L.", True],
                               ["Juan Pérez", "Constructora Sur", "Encargado", False, "Constructora Sur", True]]
                    and "Ana Cascajosa" not in str(menu2) and "Juan Pérez" not in str(menu2), str(fichas))
        empresas = pd.evaluate(OPCIONES, "asEmpresaLista0")
        anadir_asistente(pd)                                  # persona nueva
        nueva_ficha = [pd.evaluate("document.activeElement.id"), pd.evaluate("state.acta.asistentes.length")]
        pd.fill("#asNombre2", "Pedro Sanz")
        pd.select_option("#asEmpresaLista2", "__otro")
        foco_empresa = pd.evaluate("document.activeElement.id")
        pd.fill("#asEmpresa2", "Nueva Empresa S.A.")
        r.comprobar("Empresa: TEAMBIMCIVIL S.L., las de esta acta y las ya usadas (de la más reciente a la más antigua), sin repetir; «Otra…» para escribirla",
                    [e[1] for e in empresas] == ["Elige la empresa", "TEAMBIMCIVIL S.L.", "Constructora Sur", "Propiedad", "CSS Ingeniería",
                                                 "Geotecnia Andaluza", "Otra…"]
                    and nueva_ficha == ["asNombre2", 3] and foco_empresa == "asEmpresa2"
                    and pd.evaluate("state.acta.asistentes[2].empresa") == "Nueva Empresa S.A.", f"{[e[1] for e in empresas]} {nueva_ficha}")
        pd.evaluate("irAFase(5)")                              # → acuerdos
        pd.click("button:has-text('Añadir indicación')")
        pd.fill("#indTexto0", "Retirar el acopio de ferralla")
        responsables_lista = [o[1] for o in pd.evaluate(OPCIONES, "indRespLista0")]
        escribir_responsable(pd, 0, "Constructora Sur · Encargado")
        de_lista = pd.evaluate("[state.acta.indicaciones[0].responsable, document.getElementById('indResp0').classList.contains('hidden')]")
        escribir_responsable(pd, 0, "Ayuntamiento · Policía Local")
        pd.evaluate("irAFase(4)")
        pd.evaluate("irAFase(5)")
        repintado = pd.evaluate("""[document.getElementById('indRespLista0').value, document.getElementById('indResp0').value,
            document.getElementById('indResp0').classList.contains('hidden'), state.acta.indicaciones[0].responsable]""")
        r.comprobar("Responsable: desplegable con el técnico (TEAMBIMCIVIL S.L.) y los asistentes, y «Otro…» escrito que se conserva",
                    responsables_lista == ["Sin responsable", "TEAMBIMCIVIL S.L. · Director de Obra", "TEAMBIMCIVIL S.L. · Redactor de Estudio",
                                           "Constructora Sur · Encargado", "Nueva Empresa S.A.", "Otro…"]
                    and de_lista == ["Constructora Sur · Encargado", True]
                    and repintado == ["__otro", "Ayuntamiento · Policía Local", False, "Ayuntamiento · Policía Local"],
                    f"{responsables_lista} {de_lista} {repintado}")
        pd.screenshot(path=str(SALIDA / "s313-acuerdos.png"), full_page=True)

        pd.evaluate("abrirActaGuardada('d-b1')")
        pd.wait_for_function("state.acta.id === 'd-b1' && state.fase === 8")
        lectura = pd.evaluate("""[document.getElementById('expedienteLista').classList.contains('hidden'),
            document.getElementById('expediente').classList.contains('hidden'), document.getElementById('expediente').disabled,
            document.getElementById('expediente').value, document.getElementById('anadirAsistenteLista').classList.contains('hidden')]""")
        r.comprobar("Un acta emitida muestra su expediente tal cual, sin desplegables para cambiarlo ni para añadir asistentes",
                    lectura == [True, False, True, "EXP-B", True], str(lectura))
        pd.evaluate("""async () => { for (const a of await Almacen.listarActas()) await Almacen.borrarActa(a.id); }""")
        pd.evaluate("nuevaActa()")
        pd.wait_for_function("state.acta && !state.persistida && !document.getElementById('expedienteLista').classList.contains('hidden')",
                             timeout=10000)
        r.comprobar("Sin actas en el teléfono, siguen saliendo los expedientes recordados al emitir",
                    pd.evaluate(OPCIONES, "expedienteLista") == [["", "Elige el expediente"], ["EXP-RECORDADO", "EXP-RECORDADO"],
                                                                 ["__nuevo", "➕ Nuevo expediente…"]],
                    str(pd.evaluate(OPCIONES, "expedienteLista")))
        ctx_d.close()

        revision_v313(p, navegador, r, externas, errores, horizontal)

        r.comprobar("Sin peticiones a internet", not externas, "; ".join(externas[:3]))
        r.comprobar("Sin errores JS", not errores, "; ".join(errores[:3]))
        navegador.close()
    srv.shutdown()
    return r.resumen()


if __name__ == "__main__":
    sys.exit(main())
