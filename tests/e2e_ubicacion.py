"""Pruebas de la v3.1.4 en un Pixel 7 emulado: coordenadas UTM ETRS89 de la visita y de las fotos (cálculo comparado
con las coordenadas oficiales de la Red Andaluza de Posicionamiento, huso según la posición, Identificación con GPS y
con coordenadas manuales, Revisión, PDF y registro Excel; las actas emitidas antes conservan su PDF) y fotos de la
cámara o de la galería (cada botón abre su selector; las de la galería llevan su hora y solo su propio GPS).

Uso (desde la raíz del repo, con los mismos requisitos que e2e_seguimiento.py):
    python tests/e2e_ubicacion.py
"""
import datetime
import functools
import hashlib
import http.server
import io
import json
import os
import pathlib
import shutil
import subprocess
import sys
import tarfile
import threading
import zipfile

from openpyxl import load_workbook
from PIL import Image
from playwright.sync_api import sync_playwright
from pypdf import PdfReader

sys.path.insert(0, str(pathlib.Path(__file__).parent))
from comun import (PUERTO, RAIZ, SALIDA, Resultados, abrir, abrir_mis_actas, cerrar_aviso, contexto_movil,  # noqa: E402
                   esperar_fotos, firmar, identificar, servidor, texto_pdf, vigilar_red)
from comun import _ManejadorSilencioso  # noqa: E402
from generar_fotos import generar, generar_con_exif, generar_galeria_sin_gps  # noqa: E402

DATOS_V313 = pathlib.Path(__file__).parent / "datos" / "v313"     # actas emitidas con la v3.1.3 (ver su LEEME.md)
DATOS_V314 = pathlib.Path(__file__).parent / "datos" / "v314"     # ídem con la v3.1.4 (esquema 5)
PUBLICADA_V313 = "e0374d7"                                       # commit de la v3.1.3 publicada (git log origin/main)
# Estaciones de la Red Andaluza de Posicionamiento (IECA, fichas ETRS89 época 2026.0): latitud, longitud y UTM en los
# husos 29 y 30. SEV1 está en Sevilla, al oeste del meridiano 6° O (huso 29); HULV, en Huelva.
ESTACIONES = {
    "SEV1": (37.410888231, -6.002201011, {29: (765322.225, 4144674.174), 30: (234288.077, 4144686.582)}),
    "HULV": (37.280284440, -6.913543917, {29: (684972.006, 4128006.489), 30: (152998.132, 4133150.587)}),
}
GPS_SEV1 = {"latitude": 37.410888231, "longitude": -6.002201011, "accuracy": 6}
UTM_SEV1 = "Huso 29 · X 765322 · Y 4144674"          # la posición con 6 decimales, como la guarda la app
# Fases en la página: las de la v3.1.x (8) o las de FASE desde la v3.2 (con Comprobaciones).
FASES_V31 = {"FOTOS": 3, "NOTAS": 4, "ACUERDOS": 5, "FIRMAS": 6, "REVISION": 7, "EMISION": 8}


def fase(pg, nombre):
    return pg.evaluate(f"typeof FASE !== 'undefined' ? FASE.{nombre} : {FASES_V31[nombre]}")


def hasta(pg, nombre):
    destino = fase(pg, nombre)
    while pg.evaluate("state.fase") < destino:
        antes = pg.evaluate("state.fase")
        pg.click("#btnNext")
        if pg.evaluate("state.fase") == antes:
            raise AssertionError(f"No se pudo pasar de la fase {antes}: {cerrar_aviso(pg)}")


def descargar(pg, selector, ruta):
    with pg.expect_download(timeout=60000) as d:
        pg.click(selector)
    d.value.save_as(str(ruta))
    return ruta


def huellas(ruta):
    return [hashlib.sha256(pagina.get_contents().get_data()).hexdigest() for pagina in PdfReader(str(ruta)).pages]


def servir_publicada(commit, puerto):
    """Sirve en otro puerto (otro origen: otro teléfono) la versión publicada en ese commit, sacada del historial de git."""
    carpeta = SALIDA / f"publicada-{commit}"
    if carpeta.exists():
        shutil.rmtree(carpeta)
    carpeta.mkdir(parents=True)
    tar = subprocess.run(["git", "-C", str(RAIZ), "archive", "--format=tar", commit], capture_output=True, check=True).stdout
    with tarfile.open(fileobj=io.BytesIO(tar)) as t:
        t.extractall(carpeta, filter="data")
    srv = http.server.ThreadingHTTPServer(("127.0.0.1", puerto), functools.partial(_ManejadorSilencioso, directory=str(carpeta)))
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    return srv


def con_esquema(origen, destino, esquema):
    """Copia del ZIP de un acta con otro esquema (como si la hubiera exportado una versión posterior)."""
    with zipfile.ZipFile(origen) as z, zipfile.ZipFile(destino, "w") as d:
        for n in z.namelist():
            datos = z.read(n)
            if n.endswith("acta.json"):
                ficha = json.loads(datos)
                ficha["acta"]["esquema"] = esquema
                datos = json.dumps(ficha, ensure_ascii=False).encode("utf-8")
            d.writestr(n, datos)
    return destino


def elegir(pg, boton, archivos):
    """Pulsa el botón y elige los archivos en el selector que abre; devuelve (id del selector, admite varios)."""
    with pg.expect_file_chooser() as fc:
        pg.click(boton)
    selector = fc.value
    datos = (selector.element.get_attribute("id"), selector.is_multiple())
    selector.set_files([str(a) for a in archivos])
    return datos


def main():
    r = Resultados()
    SALIDA.mkdir(parents=True, exist_ok=True)
    fotos = SALIDA / "fotos"
    horizontal, vertical, _ = generar(fotos)
    con_gps, sin_gps = generar_con_exif(fotos), generar_galeria_sin_gps(fotos)
    # Foto de la galería sin datos de cámara: su hora es la del archivo (20/10/2026 09:30 en Madrid)
    sin_exif = fotos / "foto_galeria_sin_exif.jpg"
    shutil.copy(horizontal, sin_exif)
    momento = datetime.datetime(2026, 10, 20, 7, 30, tzinfo=datetime.timezone.utc).timestamp()
    os.utime(sin_exif, (momento, momento))
    # Foto con la fecha EXIF «0000:00:00 00:00:00» (la escriben algunas cámaras y apps): vale la del archivo (19/10 16:45)
    fecha_nula = fotos / "foto_fecha_nula.jpg"
    exif = Image.Exif()
    exif.get_ifd(0x8769)[0x9003] = "0000:00:00 00:00:00"
    Image.open(vertical).resize((900, 1200)).save(fecha_nula, quality=80, exif=exif)
    momento = datetime.datetime(2026, 10, 19, 14, 45, tzinfo=datetime.timezone.utc).timestamp()
    os.utime(fecha_nula, (momento, momento))
    no_imagen = fotos / "no_es_imagen.jpg"            # p. ej. una HEIC que el navegador no sabe leer
    no_imagen.write_bytes(b"esto no es una imagen")
    srv = servidor()
    externas, errores = [], []

    with sync_playwright() as p:
        navegador = p.chromium.launch()
        ctx = contexto_movil(p, navegador, gps=GPS_SEV1)
        vigilar_red(ctx, externas)
        pg = abrir(ctx, errores)
        cdp = ctx.new_cdp_session(pg)

        # ══ Cálculo de las coordenadas UTM ══
        calculadas = {(n, h): pg.evaluate("([lat, lng, h]) => { const u = utm(lat, lng, h); return [u.huso, u.x, u.y]; }",
                                          [lat, lng, h])
                      for n, (lat, lng, oficiales) in ESTACIONES.items() for h in oficiales}
        errores_mm = {f"{n} H{h}": round(max(abs(c[1] - ESTACIONES[n][2][h][0]), abs(c[2] - ESTACIONES[n][2][h][1])) * 1000, 2)
                      for (n, h), c in calculadas.items()}
        r.comprobar("UTM ETRS89: coincide con las coordenadas oficiales de las estaciones SEV1 y HULV de la RAP (IECA) en los "
                    "husos 29 y 30, con error menor de 1 mm", all(e < 1 for e in errores_mm.values()), f"error en mm: {errores_mm}")
        husos = pg.evaluate("[-6.002201011, -5.98446, -6, -6.913543917, -1.9, -7.5].map(husoDeLongitud)")
        r.comprobar("Huso por la posición: 29 al oeste del meridiano 6° O (SEV1, en la Cartuja; Huelva) y 30 al este (Sevilla "
                    "centro, Almería); el propio 6° O, en el 30", husos == [29, 30, 30, 29, 30, 29], str(husos))
        huso_fotos = pg.evaluate("""husoActa({ gps: null, sinGps: { activo: true, referencia: 'PK 1+250', lat: '', lng: '' },
            fotos: [{ gps: null }, { gps: { lat: '37.280284', lng: '-6.913544' } }, { gps: { lat: '37.389210', lng: '-5.984460' } }] })""")
        r.comprobar("Sin posición de la visita, el huso del acta es el de la primera foto con GPS", huso_fotos == 29, str(huso_fotos))

        # ══ Identificación: UTM bajo la posición GPS ══
        identificar(pg, expediente="EXP-UTM", detalle="Urbanización de la Cartuja")
        pg.click("#gpsBtn")
        pg.wait_for_selector("#gpsStatus.ok", timeout=15000)
        estado_gps = pg.inner_text("#gpsStatus")
        r.comprobar("Identificación: bajo la posición GPS, «UTM ETRS89 · Huso 29 · X … · Y …» (SEV1, al oeste del 6° O)",
                    f"UTM ETRS89 · {UTM_SEV1}" in estado_gps, estado_gps.replace("\n", " | "))
        pg.wait_for_function("state.acta.datos.actaNum === 'EXP-UTM-AV-01'", timeout=5000)

        # ══ Fotos: de la cámara o de la galería ══
        hasta(pg, "FOTOS")
        botones = pg.evaluate("""[document.getElementById('addPhotoBtn').textContent.trim(),
            document.getElementById('addGaleriaBtn').textContent.trim(),
            document.getElementById('photoInput').getAttribute('capture'), document.getElementById('galeriaInput').hasAttribute('capture'),
            document.getElementById('galeriaInput').multiple, document.getElementById('galeriaInput').accept]""")
        r.comprobar("Fotos: dos botones, «📷 Hacer foto» (la cámara) y «🖼️ Elegir de la galería» (sin cámara forzada, varias a "
                    "la vez)", botones == ["📷 Hacer foto", "🖼️ Elegir de la galería", "environment", False, True, "image/*"],
                    str(botones))
        pg.wait_for_function("gpsFotos.ultima !== null", timeout=15000)      # el GPS del teléfono ya da posición
        selector = elegir(pg, "#addGaleriaBtn", [con_gps, sin_gps, sin_exif])
        esperar_fotos(pg, 3)
        galeria = pg.evaluate("state.acta.fotos.map(f => [f.hora, f.horaFuente, f.gps])")
        r.comprobar("Galería: el botón abre su selector, con varias fotos a la vez, que entran en el orden elegido",
                    selector == ("galeriaInput", True) and len(galeria) == 3, str(selector))
        r.comprobar("Galería: cada foto con su hora (la de la cámara o, sin ella, la del archivo) y solo su propio GPS: las que no "
                    "lo traen quedan «sin GPS» aunque el teléfono tenga posición",
                    galeria == [["2026-10-21 18:05", "camara", {"lat": "37.389210", "lng": "-5.984460", "fuente": "camara"}],
                                ["2026-10-22 08:15", "camara", None], ["2026-10-20 09:30", "movil", None]], str(galeria))
        selector = elegir(pg, "#addPhotoBtn", [vertical])
        esperar_fotos(pg, 4)
        camara = pg.evaluate("state.acta.fotos[3].gps")
        r.comprobar("Cámara: «Hacer foto» abre el selector con la cámara y la foto lleva la posición del teléfono, como antes",
                    selector[0] == "photoInput" and camara == {"lat": "37.410888", "lng": "-6.002201", "acc": 6, "fuente": "movil"},
                    f"{selector} · {camara}")
        elegir(pg, "#addGaleriaBtn", [fecha_nula, no_imagen])
        aviso = pg.wait_for_selector("#alertOverlay.open", timeout=30000) and cerrar_aviso(pg)
        nula = pg.evaluate("state.acta.fotos.length === 5 && [state.acta.fotos[4].hora, state.acta.fotos[4].horaFuente]")
        estado_botones = pg.evaluate("""[document.getElementById('addPhotoBtn').textContent.trim(),
            document.getElementById('addGaleriaBtn').textContent.trim(), document.getElementById('fotosProcesando').textContent,
            document.getElementById('addPhotoBtn').disabled || document.getElementById('addGaleriaBtn').disabled,
            document.querySelector('.foto-botones').nextElementSibling.id]""")
        r.comprobar("Galería: una fecha EXIF que no existe vale como si no la trajera (la del archivo); un archivo que no se puede "
                    "leer se avisa, con la pista del formato HEIC; los botones no cambian de rótulo y «procesando» va bajo ellos",
                    nula == ["2026-10-19 16:45", "movil"] and "Fotos no añadidas" in aviso and "1 archivo(s)" in aviso
                    and "HEIC" in aviso and estado_botones == ["📷 Hacer foto", "🖼️ Elegir de la galería", "", False, "fotosProcesando"],
                    f"{nula} · {aviso[:90]} · {estado_botones}")
        metas = pg.evaluate("Array.from(document.querySelectorAll('#photosGrid .photo-meta')).map(e => e.textContent)")
        r.comprobar("En pantalla, las de la galería sin GPS dicen «sin GPS» y la de otro día lleva su fecha",
                    metas[1] == "🕒 22/10/2026 08:15 · sin GPS" and metas[2] == "🕒 20/10/2026 09:30 · sin GPS", str(metas))

        # ══ Revisión y PDF ══
        hasta(pg, "ACUERDOS")
        pg.fill("#observaciones", "Coordenadas UTM en el huso 29.")
        hasta(pg, "FIRMAS")
        firmar(cdp, pg)
        hasta(pg, "REVISION")
        revision = pg.evaluate("""Array.from(document.querySelectorAll('#reviewGeneral .review-field'))
            .map(c => [c.querySelector('.review-field-label').textContent, c.querySelector('.review-field-val').textContent])""")
        r.comprobar("Revisión: fila «UTM ETRS89» con el huso y las coordenadas de la visita, tras la localización GPS",
                    revision[-1] == ["UTM ETRS89", UTM_SEV1] and revision[-2][0] == "Localización GPS", str(revision[-2:]))
        hasta(pg, "EMISION")
        pg.wait_for_function("document.getElementById('emitPdfEstado').textContent.startsWith('PDF listo')", timeout=60000)
        ruta = descargar(pg, "#btnDescargar", SALIDA / "U-utm.pdf")
        pg.wait_for_function("state.acta.estado === 'emitida'")
        texto = texto_pdf(ruta)
        r.comprobar("PDF, Datos generales: fila «UTM ETRS89» con el huso y las coordenadas de la visita",
                    "Lat: 37.410888 | Lng: -6.002201 | Prec: ±6 m" in texto and f"UTM ETRS89 {UTM_SEV1}" in texto,
                    texto[texto.find("LOCALIZACIÓN"):][:120])
        r.comprobar("PDF, fotos: bajo la posición, la UTM en el huso del acta (también la foto al este del 6° O); las fotos sin GPS, "
                    "sin UTM",
                    "21/10/2026 18:05 · 37.389210, -5.984460 (cámara)" in texto
                    and "UTM ETRS89 · Huso 29 · X 766970 · Y 4142318" in texto
                    and "22/10/2026 08:15 · sin GPS" in texto and "20/10/2026 09:30 · sin GPS" in texto
                    and "37.410888, -6.002201 (±6 m)" in texto and f"UTM ETRS89 · {UTM_SEV1}" in texto
                    and texto.count("UTM ETRS89") == 3, texto[texto.find("REPORTAJE"):][:400])
        emitida = pg.evaluate("state.acta.id")

        # ══ Coordenadas manuales (sin cobertura) y acta sin coordenadas ══
        pg.evaluate("nuevaActa()")
        identificar(pg, expediente="EXP-MANUAL", detalle="Galería de servicios")
        pg.check("#sinGpsCheck")
        pg.fill("#sinGpsRef", "Galería de servicios, PK 0+450")
        pg.fill("#sinGpsLat", "37.280284")
        pg.fill("#sinGpsLng", "-6.913544")
        manual = pg.inner_text("#sinGpsUTM")
        pg.fill("#sinGpsLng", "abc")
        no_valida = [pg.inner_text("#sinGpsUTM"), pg.is_visible("#sinGpsUTM")]
        pg.fill("#sinGpsLng", "-6.913544")
        r.comprobar("Coordenadas manuales: su UTM sale mientras se escriben, y nada si no son válidas",
                    manual == "UTM ETRS89 · Huso 29 · X 684972 · Y 4128006" and no_valida == ["", False], f"{manual} · {no_valida}")
        pg.evaluate("guardarYa()")
        manual_id = pg.evaluate("state.acta.id")
        pg.evaluate("nuevaActa()")
        identificar(pg, expediente="EXP-SIN", detalle="Interior del túnel")
        pg.check("#sinGpsCheck")
        pg.fill("#sinGpsRef", "Túnel, PK 1+200")
        pg.evaluate("guardarYa()")
        sin_id = pg.evaluate("state.acta.id")

        # ══ Registro Excel: huso, X e Y de la visita ══
        abrir_mis_actas(pg)
        descargar(pg, "#btnRegistro", SALIDA / "U-registro.xlsx")
        cerrar_aviso(pg)
        pg.evaluate("cerrarMisActas()")
        hoja = load_workbook(SALIDA / "U-registro.xlsx")["Actas"]
        cabecera = [c.value for c in hoja[1]]
        filas = {f[5]: f[11:15] for f in hoja.iter_rows(min_row=2, values_only=True)}
        r.comprobar("Registro: columnas «Huso UTM», «UTM X» y «UTM Y» tras «Ubicación», con números; vacías sin coordenadas",
                    cabecera[11:15] == ["Ubicación", "Huso UTM", "UTM X", "UTM Y"]
                    and filas["EXP-UTM"][1:] == (29, 765322, 4144674) and filas["EXP-MANUAL"][1:] == (29, 684972, 4128006)
                    and filas["EXP-SIN"][1:] == (None, None, None), str(filas))

        # ══ Las actas emitidas antes de la v3.1.4 conservan su PDF (sin UTM) ══
        pg.set_input_files("#importInput", str(DATOS_V313 / "av01.zip"))
        pg.wait_for_selector("#alertOverlay.open", timeout=30000)
        cerrar_aviso(pg)
        pg.evaluate("cerrarMisActas()")
        with __import__("zipfile").ZipFile(DATOS_V313 / "av01.zip") as z:
            numero = json.loads(z.read(next(n for n in z.namelist() if n.endswith("acta.json"))))["acta"]["datos"]["actaNum"]
        acta_id = pg.evaluate(f"Almacen.listarActas().then(l => l.find(a => a.datos.actaNum === '{numero}').id)")
        pg.evaluate(f"abrirActaGuardada('{acta_id}')")
        pg.wait_for_function(f"state.acta.id === '{acta_id}' && state.fase === ULTIMA_FASE")
        pg.wait_for_function("document.getElementById('emitPdfEstado').textContent.startsWith('PDF listo')", timeout=60000)
        ruta = descargar(pg, "#btnDescargar", SALIDA / "U-v313.pdf")
        esperadas = json.loads((DATOS_V313 / "huellas-pdf.json").read_text(encoding="utf-8"))["av01"]
        pg.evaluate(f"irAFase({fase(pg, 'REVISION')})")
        etiquetas = pg.evaluate("Array.from(document.querySelectorAll('#reviewGeneral .review-field-label')).map(e => e.textContent)")
        r.comprobar("Acta emitida con la v3.1.3: su PDF sale idéntico, sin UTM (aunque tiene GPS), y su revisión tampoco la muestra",
                    "UTM" not in texto_pdf(ruta) and huellas(ruta) == esperadas and "Lat: " in texto_pdf(ruta)
                    and "Localización GPS" in etiquetas and "UTM ETRS89" not in etiquetas, str(etiquetas))
        r.comprobar("Los borradores y las actas emitidas con esta versión sí llevan UTM",
                    pg.evaluate(f"Promise.all(['{emitida}', '{manual_id}', '{sin_id}'].map(id => Almacen.leerActa(id)))"
                                ".then(l => l.map(a => [!emitidaAntesDe(a, UTM.desde), utmVisita(a)]))")
                    == [[True, UTM_SEV1], [True, "Huso 29 · X 684972 · Y 4128006"], [True, ""]])

        # ══ Esquema 5: un teléfono aún en la v3.1.3 no importa las actas de la v3.1.4 (rehará su PDF sin UTM) ══
        with zipfile.ZipFile(DATOS_V314 / "av01.zip") as z:
            esquema_v314 = json.loads(z.read(next(n for n in z.namelist() if n.endswith("acta.json"))))["acta"]["esquema"]
        antigua = servir_publicada(PUBLICADA_V313, PUERTO + 1)
        ctx_v = contexto_movil(p, navegador)
        vigilar_red(ctx_v, externas)
        pv = abrir(ctx_v, errores, url=f"http://localhost:{PUERTO + 1}/index.html")
        version_v = pv.evaluate("VERSION")
        pv.set_input_files("#importInput", str(DATOS_V314 / "av01.zip"))
        pv.wait_for_selector("#alertOverlay.open", timeout=30000)
        aviso_v = cerrar_aviso(pv) or ""
        en_v = pv.evaluate("Almacen.listarActas().then(l => l.length)")
        r.comprobar("Esquema 5: las actas de la v3.1.4 se exportan con el esquema 5 y un teléfono aún en la v3.1.3 no las importa",
                    esquema_v314 == 5 and pv.evaluate("ESQUEMA") == 4 and version_v == "3.1.3" and "1 no válida(s), sin importar" in aviso_v
                    and en_v == 0, f"{esquema_v314} · {version_v} · {aviso_v[:90]}")
        ctx_v.close()
        antigua.shutdown()
        # …y desde la v3.1.4, el aviso dice que hay que actualizar (para cuando llegue la versión siguiente)
        pg.set_input_files("#importInput", str(con_esquema(DATOS_V314 / "av02.zip", SALIDA / "U-esquema-siguiente.zip",
                                                           pg.evaluate("ESQUEMA") + 1)))
        pg.wait_for_selector("#alertOverlay.open", timeout=30000)
        aviso_n = cerrar_aviso(pg) or ""
        pg.evaluate("cerrarMisActas()")
        r.comprobar("Un acta de una versión posterior de la app no se importa y el aviso dice que hay que actualizar la app",
                    "1 de una versión más nueva de la app, sin importar: actualiza la app" in aviso_n and "no válida" not in aviso_n,
                    aviso_n[:160])

        r.comprobar("Sin peticiones a internet", not externas, "; ".join(externas[:3]))
        r.comprobar("Sin errores JS", not errores, "; ".join(errores[:3]))
        ctx.close()
        navegador.close()
    srv.shutdown()
    return r.resumen()


if __name__ == "__main__":
    sys.exit(main())
