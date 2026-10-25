"""Prueba de tools/reconstruir-desde-pdf.py (v3.1.2): un teléfono A con actas emitidas de todas las versiones y con
todo lo que puede llevar un acta (condiciones, GPS y sin cobertura con coordenadas, asistentes que firman y que no,
textos largos y con saltos de línea, indicaciones, seguimiento cumplido, pendiente con nuevo plazo y anulado, una
revisión, fotos con GPS del móvil, de la cámara y sin GPS, de otro día) genera sus PDF; se reconstruyen desde los PDF
y se importan en un teléfono B, que debe volver a generar los mismos PDF (texto, fotos, enlaces, formato, propiedades
y nombre; la firma, redibujada) y el mismo registro en Excel (con el seguimiento enlazado entre actas).

Uso (desde la raíz del repo; además de los requisitos de las otras pruebas: pip install pdfplumber scikit-image):
    python tests/e2e_reconstruccion.py
"""
import hashlib
import importlib.util
import io
import json
import pathlib
import subprocess
import sys

import numpy as np
import pdfplumber
from openpyxl import load_workbook
from PIL import Image
from playwright.sync_api import sync_playwright
from pypdf import PdfReader
from scipy import ndimage

sys.path.insert(0, str(pathlib.Path(__file__).parent))
from comun import (RAIZ, SALIDA, Resultados, abrir, abrir_mis_actas, cerrar_aviso, contexto_movil, marco_pdf,  # noqa: E402
                   servidor, texto_pdf, vigilar_red)
from generar_fotos import generar  # noqa: E402

HERRAMIENTA = RAIZ / "tools" / "reconstruir-desde-pdf.py"
DIR = SALIDA / "reconstruccion"

# Actas del teléfono A (se inyectan tal como las guarda la app). Firmas: trazos en un lienzo de 300 × 150.
INYECTAR = r"""async ({ fotos }) => {
  const imagen = async ruta => (await fetch(ruta)).blob();
  const blobs = {};
  for (const [clave, ruta] of Object.entries(fotos)) blobs[clave] = await imagen(ruta);
  const trazo = puntos => ({ refW: 300, refH: 150, puntos });
  const firmaTecnico = [trazo([[20, 100], [60, 30], [100, 110], [140, 25], [190, 105], [240, 40]]), trazo([[30, 130], [150, 126], [270, 132]])];
  const firmaJuan = [trazo([[40, 40], [80, 120], [120, 40], [160, 120]]), trazo([[200, 60], [260, 60]]), trazo([[230, 90]])];
  const vacia = { activo: false, referencia: '', lat: '', lng: '' };
  const sinCondiciones = { meteo: [], temperatura: '' };
  const comun = { esquema: 4, estado: 'emitida', fase: 8, seguimiento: [], seguimientoOrigen: null, seguimientoNotas: [],
                  seguimientoGuardado: null, asistentes: [], firmasAsistentes: {}, indicaciones: [], fotos: [], firma: firmaTecnico };
  const indicacionesAV01 = [
    { id: 'ind-a1-1', texto: 'Picado y reparación con mortero R4 de la viga del eje 3, con pasivado de las armaduras vistas y seguimiento con testigos de yeso durante dos semanas.', responsable: 'Constructora Sur', plazo: '2026-10-24' },
    { id: 'ind-a1-2', texto: 'Retirar el acopio de ferralla junto al vial de acceso', responsable: 'Constructora Sur', plazo: '' },
    { id: 'ind-a1-3', texto: 'Aportar el certificado del hormigón de la losa', responsable: '', plazo: '2026-10-30' },
  ];
  const asistentesAV01 = [
    { id: 'asis-a1-1', nombre: 'Juan Pérez Gómez', empresa: 'Constructora Sur', cargo: 'Jefe de Obra', firma: true },
    { id: 'asis-a1-2', nombre: 'Marta Ruiz', empresa: 'Propiedad', cargo: '', firma: true },
  ];
  const fotosAV01 = sufijo => [
    { id: `foto-a1${sufijo}-1`, clave: 'h', ancho: 1200, alto: 900, hora: '2026-10-20 09:50', horaFuente: 'movil',
      comentario: 'Fisura longitudinal en la viga del lado norte, con abertura de unos 0,3 mm y manchas de humedad en el entorno; se marca para su seguimiento.',
      gps: { lat: '37.389211', lng: '-5.984461', acc: 6, fuente: 'movil' } },
    { id: `foto-a1${sufijo}-2`, clave: 'v', ancho: 900, alto: 1200, hora: '2026-10-19 18:05', horaFuente: 'camara', comentario: '',
      gps: { lat: '37.389210', lng: '-5.984460', fuente: 'camara' } },
    { id: `foto-a1${sufijo}-3`, clave: 'c', ancho: 1200, alto: 1200, hora: '2026-10-20 09:55', horaFuente: 'movil', comentario: 'Vista general', gps: null },
  ];
  const actas = [
    { ...comun, id: 'rec-a1', creada: '2026-10-20T07:40:00.000Z', modificada: '2026-10-20T09:30:00.000Z', emitida: '2026-10-20T09:30:00.000Z',
      versionEmision: '3.1.2', revision: 0, revisiones: [{ rev: 0, fecha: '2026-10-20T09:30:00.000Z', motivo: 'Emisión inicial' }],
      sustituidaPor: 'rec-a1r1', sustituidaRev: 1,
      datos: { actaNum: 'EXP-REC-AV-01', numAuto: true, tecnico: 'Blas González', cargoTecnico: 'Director de Obra', expediente: 'EXP-REC',
               objeto: 'tajos', descripcion: 'Hormigonado de la losa de planta 1', fecha: '2026-10-20', hora: '09:45' },
      gps: { lat: '37.389210', lng: '-5.984460', acc: 8, hora: '2026-10-20T07:46:00.000Z' }, sinGps: vacia,
      condiciones: { meteo: ['nuboso', 'viento'], temperatura: '14,5' }, asistentes: asistentesAV01,
      firmasAsistentes: { 'asis-a1-1': { trazos: firmaJuan, noFirma: false }, 'asis-a1-2': { trazos: [], noFirma: true } },
      observaciones: 'Se revisa la estructura de la planta primera antes del hormigonado: encofrados, armaduras, recubrimientos, separadores y juntas de hormigonado en todo el perímetro de la losa.\nLa propiedad pide que se le avise del día del hormigonado.',
      indicaciones: indicacionesAV01, fotos: fotosAV01('') },
    { ...comun, id: 'rec-a1r1', creada: '2026-10-21T07:50:00.000Z', modificada: '2026-10-21T10:00:00.000Z', emitida: '2026-10-21T10:00:00.000Z',
      versionEmision: '3.1.2', revision: 1, corrigeA: 'rec-a1',
      revisiones: [{ rev: 0, fecha: '2026-10-20T09:30:00.000Z', motivo: 'Emisión inicial' }, { rev: 1, fecha: '2026-10-21T10:00:00.000Z', motivo: 'Corrección del detalle del objeto' }],
      datos: { actaNum: 'EXP-REC-AV-01', numAuto: false, tecnico: 'Blas González', cargoTecnico: 'Director de Obra', expediente: 'EXP-REC',
               objeto: 'tajos', descripcion: 'Hormigonado de la losa de planta 1 (zona norte)', fecha: '2026-10-20', hora: '09:45' },
      gps: { lat: '37.389210', lng: '-5.984460', acc: 8, hora: '2026-10-20T07:46:00.000Z' }, sinGps: vacia,
      condiciones: { meteo: ['nuboso', 'viento'], temperatura: '14,5' }, asistentes: asistentesAV01,
      firmasAsistentes: { 'asis-a1-1': { trazos: firmaJuan, noFirma: false }, 'asis-a1-2': { trazos: [], noFirma: true } },
      observaciones: 'Se revisa la estructura de la planta primera antes del hormigonado: encofrados, armaduras, recubrimientos, separadores y juntas de hormigonado en todo el perímetro de la losa.\nLa propiedad pide que se le avise del día del hormigonado.',
      indicaciones: indicacionesAV01.map(i => ({ ...i })), fotos: fotosAV01('r1') },
    { ...comun, id: 'rec-a2', creada: '2026-10-27T09:50:00.000Z', modificada: '2026-10-27T10:30:00.000Z', emitida: '2026-10-27T10:30:00.000Z',
      versionEmision: '3.1.2', revision: 0, revisiones: [{ rev: 0, fecha: '2026-10-27T10:30:00.000Z', motivo: 'Emisión inicial' }],
      datos: { actaNum: 'EXP-REC-AV-02', numAuto: true, tecnico: 'Ana Cascajosa', cargoTecnico: 'Inspector de Campo', expediente: 'EXP-REC',
               objeto: 'otro', descripcion: 'Reunión con la propiedad en obra', fecha: '2026-10-27', hora: '11:00' },
      gps: null, sinGps: { activo: true, referencia: 'Sótano 2, junto al ascensor', lat: '37.389000', lng: '-5.984000' },
      condiciones: sinCondiciones, observaciones: '',
      seguimiento: [
        { id: 'ind-a1-1', origen: { acta: 'EXP-REC-AV-01', rev: 1, n: 1, fecha: '2026-10-20' }, texto: indicacionesAV01[0].texto, responsable: 'Constructora Sur',
          plazo: '2026-10-24', plazoInicial: '2026-10-24', estado: 'cumplida', nuevoPlazo: '', comentario: 'Reparación ejecutada y revisada' },
        { id: 'ind-a1-2', origen: { acta: 'EXP-REC-AV-01', rev: 1, n: 2, fecha: '2026-10-20' }, texto: indicacionesAV01[1].texto, responsable: 'Constructora Sur',
          plazo: '', plazoInicial: '', estado: 'pendiente', nuevoPlazo: '2026-11-05', comentario: 'Se reitera' },
        { id: 'ind-a1-3', origen: { acta: 'EXP-REC-AV-01', rev: 1, n: 3, fecha: '2026-10-20' }, texto: indicacionesAV01[2].texto, responsable: '',
          plazo: '2026-10-30', plazoInicial: '2026-10-30', estado: 'anulada', nuevoPlazo: '', comentario: 'Ya se aportó en la visita anterior' },
      ],
      seguimientoOrigen: { expediente: 'EXP-REC', actaId: 'rec-a1r1', acta: 'EXP-REC-AV-01', rev: 1, fecha: '2026-10-20', hora: '09:45' },
      indicaciones: [{ id: 'ind-a2-1', texto: 'Limpiar la arqueta del sótano', responsable: 'Mantenimiento', plazo: '2026-11-03' }],
      fotos: [{ id: 'foto-a2-1', clave: 'h', ancho: 1200, alto: 900, hora: '2026-10-27 11:10', horaFuente: 'movil', comentario: 'Arqueta del sótano',
                gps: { lat: '37.389300', lng: '-5.984300', acc: 12, fuente: 'movil' } }] },
    { ...comun, id: 'rec-a3', creada: '2026-10-15T08:00:00.000Z', modificada: '2026-10-15T09:00:00.000Z', emitida: '2026-10-15T09:00:00.000Z',
      versionEmision: '3.1', revision: 0, revisiones: [{ rev: 0, fecha: '2026-10-15T09:00:00.000Z', motivo: 'Emisión inicial' }],
      datos: { actaNum: 'EXP-OTRO-AV-01', numAuto: true, tecnico: 'José L. Gutiérrez', cargoTecnico: 'Redactor de Estudio', expediente: 'EXP-OTRO',
               objeto: 'licitacion', descripcion: '', fecha: '2026-10-15', hora: '10:00' },
      gps: { lat: '37.401000', lng: '-5.990000', acc: 20, hora: '2026-10-15T08:01:00.000Z' }, sinGps: vacia,
      condiciones: { meteo: ['despejado'], temperatura: '-3' }, observaciones: '' },
    { ...comun, id: 'rec-a4', creada: '2026-10-22T15:00:00.000Z', modificada: '2026-10-22T16:00:00.000Z', emitida: '2026-10-22T16:00:00.000Z',
      versionEmision: '3.0', revision: 0, revisiones: [{ rev: 0, fecha: '2026-10-22T16:00:00.000Z', motivo: 'Emisión inicial' }],
      datos: { actaNum: 'V30-7', numAuto: false, tecnico: 'BGG', cargoTecnico: 'Inspector', expediente: 'EXP-V30', objeto: '',
               descripcion: 'Visita libre de la v3.0', fecha: '2026-10-22', hora: '17:00' },
      gps: null, sinGps: { activo: true, referencia: 'Interior de la nave', lat: '', lng: '' },
      condiciones: { meteo: ['lluvia'], temperatura: '' }, observaciones: 'Sin acceso a la cubierta.',
      fotos: [{ id: 'foto-a4-1', clave: 'c', ancho: 1200, alto: 1200, hora: '2026-10-22 17:05', horaFuente: 'movil', comentario: '', gps: null }] },
    { ...comun, esquema: 2, id: 'rec-a5', creada: '2026-10-21T08:00:00.000Z', modificada: '2026-10-21T08:40:00.000Z', emitida: '2026-10-21T08:40:00.000Z',
      revision: 0, revisiones: [{ rev: 0, fecha: '2026-10-21T08:40:00.000Z', motivo: 'Emisión inicial' }],
      datos: { actaNum: 'V23-1', numAuto: false, tecnico: 'BGG', cargoTecnico: 'Inspector', expediente: '168_2022 OT AFORADORES',
               descripcion: 'Implantación Equipo 3', fecha: '2026-10-21', hora: '10:15' },
      gps: { lat: '37.190805', lng: '-6.932370', acc: 46, hora: '2026-10-21T08:15:00.000Z' }, sinGps: vacia, observaciones: 'En la glorieta',
      fotos: [{ id: 'foto-a5-1', clave: 'v', ancho: 900, alto: 1200, hora: '2026-10-21 10:20', horaFuente: 'camara', comentario: 'Equipo 3',
                gps: { lat: '37.191109', lng: '-6.931233', fuente: 'camara' } }] },
    { ...comun, esquema: 2, id: 'rec-a6', creada: '2026-10-21T09:00:00.000Z', modificada: '2026-10-21T09:30:00.000Z', emitida: '2026-10-21T09:30:00.000Z',
      revision: 0, revisiones: [{ rev: 0, fecha: '2026-10-21T09:30:00.000Z', motivo: 'Emisión inicial' }],
      datos: { actaNum: 'V23-2', numAuto: false, tecnico: 'BGG', cargoTecnico: '', expediente: '168_2022 OT AFORADORES',
               descripcion: 'Revisión del túnel', fecha: '2026-10-21', hora: '11:00' },
      gps: null, sinGps: { activo: true, referencia: 'Túnel, PK 1+200', lat: '', lng: '' }, observaciones: '' },
    { ...comun, id: 'rec-a7', creada: '2026-10-24T08:00:00.000Z', modificada: '2026-10-24T09:00:00.000Z', emitida: '2026-10-24T09:00:00.000Z',
      versionEmision: '3.1.3', revision: 0, revisiones: [{ rev: 0, fecha: '2026-10-24T09:00:00.000Z', motivo: 'Emisión inicial' }],
      datos: { actaNum: 'EXP-OTRO-AV-02', numAuto: true, tecnico: 'María Quinta', cargoTecnico: 'Inspector de Campo', expediente: 'EXP-OTRO',
               objeto: 'inventario', descripcion: '', fecha: '2026-10-24', hora: '10:00' },
      gps: { lat: '37.401000', lng: '-5.990000', acc: 20, hora: '2026-10-24T08:01:00.000Z' }, sinGps: vacia,
      condiciones: sinCondiciones, observaciones: '',
      asistentes: [{ id: 'asis-a7-1', nombre: 'Sebastián Gil', empresa: 'TEAMBIMCIVIL S.L.', cargo: 'Redactor de Estudio', firma: false }] },
  ];
  for (const a of actas) {
    for (const f of a.fotos) { await Almacen.guardarFoto(f.id, a.id, blobs[f.clave]); delete f.clave; }
    if (a.esquema === 2) { delete a.condiciones; delete a.seguimiento; delete a.seguimientoOrigen; delete a.seguimientoNotas; delete a.seguimientoGuardado; }
    await Almacen.guardarActa(a);
  }
  return actas.map(a => [a.id, a.datos.actaNum, a.revision]);
}"""


def imagenes(ruta):
    fotos, firmas = [], []
    with pdfplumber.open(str(ruta)) as pdf:
        for pagina in pdf.pages:
            for im in sorted(pagina.images, key=lambda i: (round(i["top"]), i["x0"])):
                (firmas if im["srcsize"] == (750, 300) else fotos).append(im["stream"].get_rawdata())
    return fotos, firmas


def cobertura(a, b):
    """Parte de la tinta de cada firma que cae dentro de la otra engordada 2 px (1 = iguales)."""
    ga = np.asarray(Image.open(io.BytesIO(a)).convert("L")) < 140
    gb = np.asarray(Image.open(io.BytesIO(b)).convert("L")) < 140
    da, db = ndimage.binary_dilation(ga, iterations=2), ndimage.binary_dilation(gb, iterations=2)
    return min((ga & db).sum() / max(ga.sum(), 1), (gb & da).sum() / max(gb.sum(), 1))


def enlaces(ruta):
    return [o.get_object()["/A"].get_object().get("/URI") for pagina in PdfReader(str(ruta)).pages
            for o in (pagina.get("/Annots") or []) if "/A" in o.get_object()]


def hojas(ruta):
    libro = load_workbook(ruta)
    actas = [list(f) for f in libro["Actas"].iter_rows(values_only=True)]
    i = actas[0].index("Emitida")
    return [f[:i] + f[i + 1:] for f in actas], [list(f) for f in libro["Indicaciones"].iter_rows(values_only=True)]


def descargar_pdf(pg, acta_id, ruta):
    pg.evaluate(f"abrirActaGuardada('{acta_id}')")
    pg.wait_for_function(f"state.acta.id === '{acta_id}' && state.fase === 8")
    pg.wait_for_function("document.getElementById('emitPdfEstado').textContent.startsWith('PDF listo')", timeout=60000)
    with pg.expect_download(timeout=60000) as d:
        pg.click("#btnDescargar")
    d.value.save_as(str(ruta))
    cerrar_aviso(pg)
    return d.value.suggested_filename


def registro(pg, ruta):
    abrir_mis_actas(pg)
    with pg.expect_download(timeout=60000) as d:
        pg.click("#btnRegistro")
    d.value.save_as(str(ruta))
    cerrar_aviso(pg)
    pg.evaluate("cerrarMisActas()")


def main():
    r = Resultados()
    (DIR / "originales").mkdir(parents=True, exist_ok=True)
    (DIR / "regenerados").mkdir(parents=True, exist_ok=True)
    horizontal, vertical, cuadrada = generar(SALIDA / "fotos")
    fotos = {}
    for clave, ruta, tam in (("h", horizontal, (1200, 900)), ("v", vertical, (900, 1200)), ("c", cuadrada, (1200, 1200))):
        destino = DIR / f"foto-{clave}.jpg"
        Image.open(ruta).resize(tam).save(destino, quality=75)
        fotos[clave] = f"tests/salida/reconstruccion/{destino.name}"
    herramienta = importlib.util.spec_from_file_location("reconstruir", HERRAMIENTA)
    modulo = importlib.util.module_from_spec(herramienta)
    herramienta.loader.exec_module(modulo)
    srv = servidor()
    externas, errores = [], []

    with sync_playwright() as p:
        navegador = p.chromium.launch()
        # ══ Teléfono A: las actas originales y sus PDF ══
        ctx = contexto_movil(p, navegador)
        vigilar_red(ctx, externas)
        pa = abrir(ctx, errores)
        listas = pa.evaluate("[METEO.map(m => [m.nombre, m.clave]), OBJETOS.filter(o => o.clave !== 'otro').map(o => [o.nombre, o.clave])]")
        r.comprobar("Las listas de condiciones y objetos de la herramienta coinciden con las de la app",
                    listas == [list(map(list, modulo.METEO.items())), list(map(list, modulo.OBJETOS.items()))])
        inyectadas = pa.evaluate(INYECTAR, {"fotos": fotos})
        originales = {}
        for acta_id, numero, rev in inyectadas:
            ruta = DIR / "originales" / f"{acta_id}.pdf"
            originales[(numero, rev)] = (acta_id, ruta, descargar_pdf(pa, acta_id, ruta))
        registro(pa, DIR / "registro-A.xlsx")
        textos_a = {k: texto_pdf(v[1]) for k, v in originales.items()}
        r.comprobar("Teléfono A: 8 actas de la v2.3 a la v3.1.3 con sus PDF (formatos «Acta de Campo» y anterior, con y sin enlaces)",
                    len(originales) == 8 and sum("ACTA DE CAMPO" in t for t in textos_a.values()) == 4
                    and "María Quinta TEAMBIMCIVIL S.L. Inspector de Campo" in textos_a[("EXP-OTRO-AV-02", 0)]
                    and "Blas González TEAMBIMCIVIL Director de Obra" in textos_a[("EXP-REC-AV-01", 1)]
                    and "6 · FIRMAS" in textos_a[("V23-2", 0)] and "5 · FIRMAS" in textos_a[("EXP-OTRO-AV-01", 0)]
                    and len(enlaces(originales[("EXP-OTRO-AV-01", 0)][1])) == 1 and not enlaces(originales[("V23-1", 0)][1]),
                    "; ".join(f"{k[0]} Rev. {k[1]}" for k in originales))
        ctx.close()

        # ══ Reconstrucción desde los PDF ══
        salida = subprocess.run([sys.executable, str(HERRAMIENTA), str(DIR / "zip")] + [str(v[1]) for v in originales.values()],
                                capture_output=True, text=True, timeout=600)
        print(salida.stdout.strip())
        r.comprobar("La herramienta reconstruye los 8 PDF sin fallos", salida.returncode == 0 and "8 acta(s) reconstruida(s)" in salida.stdout,
                    salida.stderr.strip()[-300:])
        with __import__("zipfile").ZipFile(DIR / "zip" / "Actas_reconstruidas.zip") as z:
            fichas = {json.loads(z.read(n))["acta"]["datos"]["actaNum"] + f" Rev. {json.loads(z.read(n))['acta']['revision']}": json.loads(z.read(n))["acta"]
                      for n in z.namelist() if n.endswith("acta.json")}
        versiones = {k: f.get("versionEmision") for k, f in fichas.items()}
        r.comprobar("Versión de emisión deducida del contenido (no de la app que generó el PDF): «TEAMBIMCIVIL S.L.» desde la v3.1.3",
                    versiones["EXP-REC-AV-01 Rev. 0"] == versiones["EXP-REC-AV-02 Rev. 0"] == "3.1.2"
                    and versiones["EXP-OTRO-AV-02 Rev. 0"] == "3.1.3"
                    and versiones["EXP-OTRO-AV-01 Rev. 0"] == "3.1" and versiones["V30-7 Rev. 0"] in ("3.0", "3.1")
                    and versiones["V23-1 Rev. 0"] is None and versiones["V23-2 Rev. 0"] is None, str(versiones))
        a2 = fichas["EXP-REC-AV-02 Rev. 0"]
        r1 = fichas["EXP-REC-AV-01 Rev. 1"]
        r.comprobar("Seguimiento enlazado con las indicaciones de la Rev. 1 del lote, y la Rev. 0 sustituida por la Rev. 1",
                    [s["id"] for s in a2["seguimiento"]] == [i["id"] for i in r1["indicaciones"]]
                    and a2["seguimientoOrigen"]["actaId"] == r1["id"] and r1.get("corrigeA") == fichas["EXP-REC-AV-01 Rev. 0"]["id"]
                    and fichas["EXP-REC-AV-01 Rev. 0"].get("sustituidaPor") == r1["id"], str([s["id"] for s in a2["seguimiento"]]))

        # ══ Teléfono B: importa lo reconstruido y vuelve a generar cada PDF ══
        ctx = contexto_movil(p, navegador)
        vigilar_red(ctx, externas)
        pb = abrir(ctx, errores)
        pb.set_input_files("#importInput", str(DIR / "zip" / "Actas_reconstruidas.zip"))
        pb.wait_for_selector("#alertOverlay.open", timeout=60000)
        aviso = cerrar_aviso(pb) or ""
        r.comprobar("Teléfono B: importa las 8 actas reconstruidas", "8 acta(s) nueva(s)" in aviso and "no válida" not in aviso, aviso)
        en_b = {(a["datos"]["actaNum"], a["revision"]): a for a in pb.evaluate("Almacen.listarActas()")}
        for clave, (_, ruta_a, nombre_a) in originales.items():
            acta_b = en_b[clave]
            ruta_b = DIR / "regenerados" / ruta_a.name
            nombre_b = descargar_pdf(pb, acta_b["id"], ruta_b)
            fotos_a, firmas_a = imagenes(ruta_a)
            fotos_b, firmas_b = imagenes(ruta_b)
            meta_a, meta_b = PdfReader(str(ruta_a)).metadata, PdfReader(str(ruta_b)).metadata
            iguales = texto_pdf(ruta_a) == texto_pdf(ruta_b)
            if not iguales:
                import difflib
                print("\n".join(list(difflib.unified_diff(texto_pdf(ruta_a).split(" · "), texto_pdf(ruta_b).split(" · "), lineterm="", n=0))[:12]))
            cubre = [round(float(cobertura(x, y)), 3) for x, y in zip(firmas_a, firmas_b)]
            r.comprobar(f"{clave[0]} Rev. {clave[1]}: mismo texto, fotos, enlaces, marco, propiedades y nombre; firmas ≥ 97 %",
                        iguales and [hashlib.sha256(f).hexdigest() for f in fotos_a] == [hashlib.sha256(f).hexdigest() for f in fotos_b]
                        and enlaces(ruta_a) == enlaces(ruta_b) and marco_pdf(ruta_a) == marco_pdf(ruta_b)
                        and all(meta_a.get(k) == meta_b.get(k) for k in ("/Title", "/Subject", "/Author", "/Keywords"))
                        and nombre_a == nombre_b and len(firmas_a) == len(firmas_b) and all(c >= 0.97 for c in cubre),
                        f"{len(fotos_b)} fotos · {len(enlaces(ruta_b))} enlaces · firmas {cubre}")
        registro(pb, DIR / "registro-B.xlsx")
        actas_a, ind_a = hojas(DIR / "registro-A.xlsx")
        actas_b, ind_b = hojas(DIR / "registro-B.xlsx")
        r.comprobar("Registro en Excel igual en los dos teléfonos (hoja Actas, salvo la hora de emisión)",
                    sorted(actas_a[1:], key=str) == sorted(actas_b[1:], key=str) and actas_a[0] == actas_b[0])
        r.comprobar("…y la hoja Indicaciones igual: el seguimiento reconstruido actualiza el estado de las indicaciones de la AV-01",
                    ind_a == ind_b and len(ind_b) == 5 and [f[8] for f in ind_b[1:]] == ["Cumplida", "Pendiente", "Anulada", "Pendiente"],
                    str([f[8] for f in ind_b[1:]]))
        ctx.close()

        r.comprobar("Sin peticiones a internet", not externas, "; ".join(externas[:3]))
        r.comprobar("Sin errores JS", not errores, "; ".join(errores[:3]))
        navegador.close()
    srv.shutdown()
    return r.resumen()


if __name__ == "__main__":
    sys.exit(main())
