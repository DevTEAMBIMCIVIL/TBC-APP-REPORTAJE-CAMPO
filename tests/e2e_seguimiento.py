"""Pruebas de la v3.0 en un Pixel 7 emulado: condiciones de la visita, seguimiento de las indicaciones entre
visitas (estados, motivo obligatorio, nuevo plazo, vencidas, cambio de fuente), PDF con la sección de seguimiento
y los enlaces al mapa, revisiones, registro de actas en Excel (.xlsx), migración de las actas de la v2.3,
exportación e importación con seguimiento y precarga del service worker con «?v=<VERSION>».

Uso (desde la raíz del repo, con los mismos requisitos que e2e_acta.py y además openpyxl):
    pip install openpyxl
    python tests/e2e_seguimiento.py
"""
import copy
import datetime
import json
import pathlib
import re
import sys
import zipfile

from openpyxl import load_workbook
from pypdf import PdfReader
from playwright.sync_api import sync_playwright

sys.path.insert(0, str(pathlib.Path(__file__).parent))
from comun import (MARCO_NEGRO, SALIDA, TECNICO, Resultados, abrir, abrir_mis_actas, avanzar_hasta,  # noqa: E402
                   cerrar_aviso, contexto_movil, esperar_fotos, firmar, marco_pdf, servidor, texto_pdf, vigilar_red)
from comun import escribir_expediente, escribir_responsable  # noqa: E402
from generar_fotos import generar  # noqa: E402

EXP = "EXP-SEG"
MAPA = "https://www.google.com/maps/search/?api=1&query=37.389210,-5.984460"
ASISTENTE = {"nombre": "Juan Pérez Gómez", "empresa": "Constructora Sur", "cargo": "Jefe de obra"}
INDICACIONES_AV01 = [   # texto, responsable, plazo (la 1.ª vence antes de la visita siguiente, el 15/10)
    ("Picado y reparación con mortero R4 de la viga del eje 3.", "Constructora Sur · Jefe de obra", "2026-10-12"),
    ("Aportar el último ensayo de esclerometría.", "Propiedad · Promotora", "2026-10-30"),
    ("Retirar los acopios de la zona de paso.", "Constructora Sur · Jefe de obra", ""),
]
MOTIVO = "La zona ya está despejada; no procede"
SECCIONES_V23 = ["1 · DATOS GENERALES", "2 · ASISTENTES", "3 · OBSERVACIONES GENERALES", "4 · INDICACIONES",
                 "5 · REPORTAJE FOTOGRÁFICO (1)", "6 · FIRMAS", "7 · CONTROL DE REVISIONES"]
SECCIONES_V30 = ["1 · DATOS GENERALES", "2 · ASISTENTES", "3 · OBSERVACIONES GENERALES",
                 "4 · SEGUIMIENTO DE INDICACIONES ANTERIORES", "5 · INDICACIONES", "6 · REPORTAJE FOTOGRÁFICO (1)",
                 "7 · FIRMAS", "8 · CONTROL DE REVISIONES"]
VERDE, AMBAR = "0.118 0.478 0.239 rg", "0.541 0.353 0. rg"    # colores de «Cumplida» y «Pendiente» (jsPDF escribe «0.»)

# Tarjetas del seguimiento en pantalla: origen, plazo, vencida, texto y estado marcado.
TARJETAS = """Array.from(document.querySelectorAll('#seguimientoBloque .seg-card')).map(t => ({
  origen: t.querySelector('.seg-origen').textContent, plazo: t.querySelector('.seg-plazo').textContent,
  vencida: t.classList.contains('vencida'), texto: t.querySelector('.seg-texto').textContent,
  estado: (t.querySelector('.seg-estado[aria-pressed="true"]') || { dataset: {} }).dataset.estado }))"""
PASTILLAS = "Array.from(document.querySelectorAll('#seguimientoBloque .seg-pill')).map(p => p.textContent)"


def descargar(pg, selector, ruta):
    with pg.expect_download(timeout=60000) as d:
        pg.click(selector)
    d.value.save_as(str(ruta))
    return d.value.suggested_filename


def enlaces(ruta):
    """URL de los enlaces del PDF, en orden: [(página, url)]."""
    salida = []
    for n, pagina in enumerate(PdfReader(str(ruta)).pages, 1):
        for anotacion in pagina.get("/Annots") or []:
            o = anotacion.get_object()
            if o.get("/Subtype") == "/Link" and "/A" in o:
                salida.append((n, o["/A"].get_object().get("/URI")))
    return salida


def contenido_paginas(ruta):
    return [pagina.get_contents().get_data().decode("latin-1") for pagina in PdfReader(str(ruta)).pages]


def firmar_todos(cdp, pg):
    firmar(cdp, pg)
    firmar(cdp, pg, lienzo=pg.locator("canvas.firma-lienzo").first)


def esperar_aviso(pg, espera=30000):
    pg.wait_for_selector("#alertOverlay.open", timeout=espera)
    return cerrar_aviso(pg) or ""


def leer(pg, acta_id):
    return pg.evaluate("id => Almacen.leerActa(id)", acta_id)


def nueva_visita(pg, acta_id, numero, fecha):
    """«Mis actas → Nueva visita» desde un acta: fecha de la visita y GPS."""
    abrir_mis_actas(pg)
    pg.click(f".acta-card[data-acta='{acta_id}'] button:has-text('Nueva visita')")
    pg.wait_for_function(f"state.acta.datos.actaNum === '{numero}'", timeout=10000)
    pg.fill("#fechaVisita", fecha)
    pg.click("#gpsBtn")
    pg.wait_for_selector("#gpsStatus.ok", timeout=15000)


def hasta_acuerdos(pg, foto):
    """Desde la identificación: asistentes → una foto → notas → acuerdos, con el seguimiento ya cargado."""
    pg.click("#btnNext")                       # → asistentes
    pg.click("#btnNext")                       # → fotos
    pg.set_input_files("#photoInput", [foto])
    esperar_fotos(pg, 1)
    pg.click("#btnNext")                       # → notas
    pg.locator("textarea.comment-textarea").nth(0).fill("Vista general de la zona revisada")
    pg.click("#btnNext")                       # → acuerdos
    pg.wait_for_function("state.fase === 5")
    pg.evaluate("cargarSeguimiento()")         # espera a que termine la carga del seguimiento


def main():
    r = Resultados()
    SALIDA.mkdir(parents=True, exist_ok=True)
    horizontal, vertical, cuadrada = (str(f) for f in generar(SALIDA / "fotos"))
    hoy = datetime.date.today().isoformat()
    srv = servidor()
    externas, errores = [], []

    with sync_playwright() as p:
        navegador = p.chromium.launch()
        ctx = contexto_movil(p, navegador)
        vigilar_red(ctx, externas)
        pg = abrir(ctx, errores)
        cdp = ctx.new_cdp_session(pg)
        version, esquema = pg.evaluate("[VERSION, ESQUEMA]")

        # ══ AV-01 · Condiciones de la visita (mejora 5) ══
        escribir_expediente(pg, EXP)
        pg.wait_for_function(f"state.acta.datos.actaNum === '{EXP}-AV-01'", timeout=5000)
        pg.select_option("#tecnico", TECNICO)
        pg.select_option("#cargoTecnico", "Director de Obra")
        pg.select_option("#objeto", "tajos")
        pg.fill("#descripcion", "Revisión de la estructura de planta primera")
        pg.fill("#fechaVisita", "2026-10-08")
        pg.fill("#horaVisita", "10:30")
        chips = pg.evaluate("""Array.from(document.querySelectorAll('#meteoGrid [data-meteo]'))
            .map(b => [b.dataset.meteo, b.textContent, b.getAttribute('aria-pressed')])""")
        r.comprobar("Condiciones: 7 opciones de tiempo, ninguna marcada al empezar",
                    [c[0] for c in chips] == ["despejado", "nuboso", "lluvia", "viento", "niebla", "calor", "frio"]
                    and all(c[2] == "false" for c in chips), ", ".join(c[1] for c in chips))
        for clave in ("viento", "despejado", "nuboso", "despejado"):
            pg.click(f"#meteoGrid [data-meteo='{clave}']")
        marcadas = pg.evaluate("state.acta.condiciones.meteo")
        pulsadas = pg.evaluate("Array.from(document.querySelectorAll('#meteoGrid [aria-pressed=\"true\"]')).map(b => b.dataset.meteo)")
        r.comprobar("Condiciones: varias a la vez, en el orden de la lista; al pulsar otra vez se desmarca",
                    marcadas == ["nuboso", "viento"] and pulsadas == ["nuboso", "viento"], str(marcadas))
        pg.click("#gpsBtn")
        pg.wait_for_selector("#gpsStatus.ok", timeout=15000)
        avisos = []
        for valor in ("abc", "70", "-31", "14,55"):
            pg.fill("#temperatura", valor)
            pg.click("#btnNext")
            avisos.append(cerrar_aviso(pg) or "")
        r.comprobar("Temperatura no válida (texto, fuera de -30…55 °C, dos decimales): se avisa y no se avanza",
                    all("Temperatura no válida" in a for a in avisos) and pg.evaluate("state.fase") == 1,
                    " | ".join(a[:28] for a in avisos))
        pg.fill("#temperatura", "3")
        pg.click("#temperaturaSigno")
        bajo_cero = [pg.input_value("#temperatura"), pg.evaluate("textoCondiciones(state.acta)")]
        pg.click("#temperaturaSigno")
        r.comprobar("«±» pone y quita el signo menos (el teclado numérico del iPhone no lo tiene)",
                    bajo_cero == ["-3", "Nuboso, viento · -3 °C"] and pg.input_value("#temperatura") == "3", str(bajo_cero))
        pg.fill("#temperatura", "14,5")
        pg.screenshot(path=str(SALIDA / "s-condiciones.png"), full_page=True)
        pg.click("#btnNext")
        condiciones = pg.evaluate("textoCondiciones(state.acta)")
        r.comprobar("Temperatura válida: se avanza y el texto es «Nuboso, viento · 14,5 °C»",
                    pg.evaluate("state.fase") == 2 and condiciones == "Nuboso, viento · 14,5 °C", condiciones)
        sin_gps = pg.evaluate("""(() => { const a = actaNueva();
            a.sinGps = { activo: true, referencia: 'Sótano 2', lat: '', lng: '' }; const sinCoordenadas = enlaceMapa(a);
            a.sinGps.lat = '37,38921'; a.sinGps.lng = '-5.98446'; return [sinCoordenadas, enlaceMapa(a)]; })()""")
        r.comprobar("Sin cobertura GPS: enlace al mapa solo si hay coordenadas manuales", sin_gps == [None, MAPA], str(sin_gps))
        # Dos visitas el mismo día: la fuente es la anterior en hora, no la de la tarde; el expediente, sin distinguir mayúsculas.
        mismo_dia = pg.evaluate("""(() => {
            const emitida = (numero, hora) => Object.assign(actaNueva(), { estado: 'emitida', emitida: '2026-10-23T18:00:00.000Z',
              datos: { ...actaNueva().datos, actaNum: numero, expediente: 'EXP-X', fecha: '2026-10-23', hora } });
            const tarde = emitida('EXP-X-AV-01', '17:00'), manana = emitida('EXP-X-AV-02', '09:00');
            const borrador = actaNueva();
            borrador.datos = { ...borrador.datos, actaNum: 'EXP-X-AV-03', expediente: ' exp-x ', fecha: '2026-10-23', hora: '12:00' };
            const deMediodia = fuenteSeguimiento([tarde, manana], borrador);
            borrador.datos.hora = '18:30';
            return [deMediodia && deMediodia.datos.actaNum, fuenteSeguimiento([tarde, manana], borrador).datos.actaNum];
        })()""")
        r.comprobar("Dos visitas el mismo día: la fuente es la anterior en hora (expediente sin distinguir mayúsculas)",
                    mismo_dia == ["EXP-X-AV-02", "EXP-X-AV-01"], str(mismo_dia))

        pg.evaluate("d => anadirAsistente(d)", ASISTENTE)
        pg.click("#btnNext")                               # → fotos
        pg.set_input_files("#photoInput", [horizontal])
        esperar_fotos(pg, 1)
        pg.click("#btnNext")                               # → notas
        pg.locator("textarea.comment-textarea").nth(0).fill("Fisura en la viga del eje 3")
        pg.click("#btnNext")                               # → acuerdos
        pg.evaluate("cargarSeguimiento()")
        r.comprobar("Primera visita del expediente: sin seguimiento y con el título «Indicaciones»",
                    pg.is_hidden("#seguimientoBloque") and pg.text_content("#indTitulo") == "Indicaciones"
                    and pg.evaluate("state.acta.seguimiento.length === 0 && state.acta.seguimientoOrigen === null"))
        for i, (texto, responsable, plazo) in enumerate(INDICACIONES_AV01):
            pg.click("button:has-text('Añadir indicación')")
            pg.fill(f"#indTexto{i}", texto)
            escribir_responsable(pg, i, responsable)
            if plazo:
                pg.fill(f"#indPlazo{i}", plazo)
        pg.click("#btnNext")                               # → firmas
        firmar_todos(cdp, pg)
        pg.click("#btnNext")                               # → revisión
        general = pg.inner_text("#reviewGeneral")
        r.comprobar("Revisión: las condiciones en Datos generales", "Condiciones" in general and condiciones in general)
        pg.click("#btnNext")                               # → emisión
        ruta1 = SALIDA / "S-AV01.pdf"
        descargar(pg, "#btnDescargar", ruta1)
        pg.wait_for_function("state.acta.estado === 'emitida'")
        av01 = pg.evaluate("state.acta.id")
        texto1 = texto_pdf(ruta1)
        r.comprobar("PDF sin seguimiento: las mismas 7 secciones que en la v2.3",
                    all(s in texto1 for s in SECCIONES_V23) and "SEGUIMIENTO" not in texto1,
                    ", ".join(s for s in SECCIONES_V23 if s not in texto1))
        r.comprobar("PDF: fila «Condiciones» en Datos generales", "CONDICIONES" in texto1 and condiciones in texto1)
        enlaces1 = enlaces(ruta1)
        r.comprobar("PDF: «Ver en el mapa» tras la ubicación y enlace en las coordenadas de la foto (Google Maps)",
                    "Prec: ±8 m · Ver en el mapa" in texto1 and [u for _, u in enlaces1] == [MAPA, MAPA], str(enlaces1))
        r.comprobar("Al emitir se guarda la versión de la app (el PDF de las emitidas antes se rehace sin enlaces)",
                    pg.evaluate("state.acta.versionEmision") == version)

        # ══ AV-02 · «Nueva visita»: seguimiento de las indicaciones de la AV-01 (mejora 1) ══
        nueva_visita(pg, av01, f"{EXP}-AV-02", "2026-10-15")
        r.comprobar("«Nueva visita» no copia las condiciones de la visita anterior",
                    pg.evaluate("state.acta.condiciones") == {"meteo": [], "temperatura": ""}
                    and pg.evaluate("document.querySelectorAll('#meteoGrid [aria-pressed=\"true\"]').length") == 0)
        pg.click("#meteoGrid [data-meteo='lluvia']")
        pg.fill("#temperatura", "9")
        hasta_acuerdos(pg, vertical)
        av02 = pg.evaluate("state.acta.id")
        t = pg.evaluate(TARJETAS)
        r.comprobar("Al entrar en Acuerdos se cargan las 3 indicaciones de la AV-01, todas pendientes",
                    [x["origen"] for x in t] == ["AV-01 · 1", "AV-01 · 2", "AV-01 · 3"]
                    and all(x["estado"] == "pendiente" for x in t)
                    and pg.evaluate("state.acta.seguimientoOrigen.actaId") == av01, str([x["origen"] for x in t]))
        r.comprobar("«VENCIDA»: pendiente con el plazo anterior a la fecha de la visita",
                    len(t) == 3 and t[0]["vencida"] and t[0]["plazo"] == "Plazo 12/10/2026 · VENCIDA"
                    and not t[1]["vencida"] and t[2]["plazo"] == "Sin plazo", " | ".join(x["plazo"] for x in t))
        cabecera = pg.inner_text("#seguimientoBloque .seg-cabecera")
        pastillas = pg.evaluate(PASTILLAS)
        r.comprobar("Cabecera del seguimiento (origen y resumen) y título «Nuevas indicaciones»",
                    "INDICACIONES DE VISITAS ANTERIORES (3)" in cabecera.upper() and "última: AV-01 del 08/10/2026" in cabecera
                    and pastillas == ["3 pendientes", "1 vencida"] and pg.text_content("#indTitulo") == "Nuevas indicaciones",
                    f"{cabecera.replace(chr(10), ' ')} | {pastillas}")

        pg.click("[data-seguimiento='0'] [data-estado='cumplida']")
        pg.fill("#segCom0", "Reparación ejecutada")
        pg.fill("#segPlazo1", "2026-10-10")
        pg.fill("#segCom1", "Se reitera a la propiedad")
        pg.click("[data-seguimiento='2'] [data-estado='anulada']")
        foco = pg.evaluate("document.activeElement.id")
        pg.click("#btnNext")
        aviso = cerrar_aviso(pg) or ""
        r.comprobar("«Anulada» pide el motivo (lo enfoca) y sin él no se avanza",
                    "Falta el motivo" in aviso and "AV-01 · 3" in aviso and foco == "segCom2" and pg.evaluate("state.fase") == 5,
                    f"{aviso} · foco {foco}")
        pg.fill("#segCom2", MOTIVO)
        pg.click("#btnNext")
        aviso = cerrar_aviso(pg) or ""
        r.comprobar("El nuevo plazo no puede ser anterior a la fecha de la visita",
                    "Nuevo plazo no válido" in aviso and "15/10/2026" in aviso and pg.evaluate("state.fase") == 5, aviso)
        pg.fill("#segPlazo1", "2026-11-05")
        pg.click("button:has-text('Añadir indicación')")
        pg.fill("#indTexto0", "Colocar testigos de yeso en la fisura.")
        escribir_responsable(pg, 0, "Constructora Sur · Jefe de obra")
        pg.fill("#indPlazo0", "2026-10-20")
        seguimiento = pg.evaluate("state.acta.seguimiento.map(s => [s.estado, s.nuevoPlazo, s.comentario])")
        pastillas = pg.evaluate(PASTILLAS)
        r.comprobar("Estados anotados: cumplida, pendiente con nuevo plazo y anulada con motivo",
                    seguimiento == [["cumplida", "", "Reparación ejecutada"], ["pendiente", "2026-11-05", "Se reitera a la propiedad"],
                                    ["anulada", "", MOTIVO]] and pastillas == ["1 cumplida", "1 pendiente", "1 anulada"],
                    f"{seguimiento} | {pastillas}")
        pg.screenshot(path=str(SALIDA / "s-seguimiento.png"), full_page=True)
        pg.click("#btnNext")                               # → firmas
        firmar_todos(cdp, pg)
        pg.click("#btnNext")                               # → revisión
        acuerdos = pg.inner_text("#reviewAcuerdos")
        r.comprobar("Revisión: el seguimiento antes de las observaciones, con su estado y comentario",
                    0 <= acuerdos.find("SEGUIMIENTO DE INDICACIONES ANTERIORES (3)") < acuerdos.find("OBSERVACIONES")
                    and "Cumplida · Reparación ejecutada" in acuerdos
                    and "Pendiente · nuevo plazo 05/11/2026 · Se reitera a la propiedad" in acuerdos
                    and f"Anulada · {MOTIVO}" in acuerdos and "NUEVAS INDICACIONES (1)" in acuerdos, acuerdos[:160])
        pg.screenshot(path=str(SALIDA / "s-revision.png"), full_page=True)
        pg.click("#btnNext")                               # → emisión
        mensaje = pg.evaluate("textoCompartir()")
        r.comprobar("Texto que acompaña al PDF: línea del seguimiento",
                    "Seguimiento: 1 cumplida · 1 pendiente · 1 anulada" in mensaje, mensaje.replace("\n", " | "))
        ruta2 = SALIDA / "S-AV02.pdf"
        descargar(pg, "#btnDescargar", ruta2)
        pg.wait_for_function("state.acta.estado === 'emitida'")

        # ── PDF de la AV-02 (mejoras 1 y 4) ──
        texto2 = texto_pdf(ruta2)
        r.comprobar("PDF con seguimiento: sección 4 y numeración ajustada (8 secciones)",
                    all(s in texto2 for s in SECCIONES_V30), ", ".join(s for s in SECCIONES_V30 if s not in texto2))
        r.comprobar("PDF: resumen del seguimiento",
                    "3 indicaciones de actas anteriores: 1 cumplida · 1 pendiente · 1 anulada" in texto2)
        r.comprobar("PDF: tabla del seguimiento (origen, estado, nuevo plazo y comentario)",
                    all(s in texto2 for s in ("AV-01 · 1", "Cumplida", "Reparación ejecutada", "30/10/2026 · nuevo 05/11/2026",
                                              "Pendiente", "Anulada", MOTIVO, "Colocar testigos de yeso en la fisura.")))
        paginas2 = contenido_paginas(ruta2)
        r.comprobar("PDF: estados en color (verde «Cumplida», ámbar «Pendiente»)",
                    any(VERDE in c and AMBAR in c for c in paginas2))
        r.comprobar("PDF: condiciones de la AV-02 y enlaces a Google Maps (ubicación y foto)",
                    "Lluvia · 9 °C" in texto2 and [u for _, u in enlaces(ruta2)] == [MAPA, MAPA])

        # ══ AV-03: hereda solo lo que la AV-02 deja abierto ══
        nueva_visita(pg, av02, f"{EXP}-AV-03", "2026-10-23")
        hasta_acuerdos(pg, cuadrada)
        av03 = pg.evaluate("state.acta.id")
        t = pg.evaluate(TARJETAS)
        r.comprobar("AV-03: solo la pendiente que deja la AV-02 y la indicación nueva de la AV-02",
                    [x["origen"] for x in t] == ["AV-01 · 2", "AV-02 · 1"], str([x["origen"] for x in t]))
        r.comprobar("AV-03: el plazo vigente es el nuevo (con el inicial) y la de la AV-02 ya vence",
                    len(t) == 2 and t[0]["plazo"] == "Plazo 05/11/2026 (inicial 30/10/2026)" and not t[0]["vencida"]
                    and t[1]["vencida"] and t[1]["plazo"] == "Plazo 20/10/2026 · VENCIDA", " | ".join(x["plazo"] for x in t))
        pg.click("[data-seguimiento='0'] [data-estado='cumplida']")
        pg.fill("#segCom0", "Ensayo entregado")
        pg.fill("#segPlazo1", "2026-10-30")
        pg.fill("#segCom1", "Testigos colocados; falta la foto a los 15 días")
        anotado = pg.evaluate("state.acta.seguimiento.map(s => [s.id, s.estado, s.nuevoPlazo, s.comentario])")

        def cambiar_fecha(fecha, origen_id):
            """Corrige la fecha de la visita y vuelve a Acuerdos: devuelve el aviso de «Seguimiento actualizado»."""
            pg.evaluate("irAFase(1)")
            pg.fill("#fechaVisita", fecha)
            pg.evaluate("irAFase(5)")
            if origen_id:
                pg.wait_for_function(f"state.acta.seguimientoOrigen && state.acta.seguimientoOrigen.actaId === '{origen_id}'")
            else:
                pg.wait_for_function("state.acta.seguimientoOrigen === null")
            return esperar_aviso(pg)

        aviso = cambiar_fecha("2026-10-14", av01)          # antes de la AV-02: la fuente pasa a ser la AV-01
        t = pg.evaluate(TARJETAS)
        r.comprobar("Si cambia la fuente (otra fecha), se rehace la lista, se conserva lo anotado y se avisa",
                    [(x["origen"], x["estado"]) for x in t] == [("AV-01 · 1", "pendiente"), ("AV-01 · 2", "cumplida"), ("AV-01 · 3", "pendiente")]
                    and pg.evaluate("state.acta.seguimiento[1].comentario") == "Ensayo entregado"
                    and "Seguimiento actualizado" in aviso and f"{EXP}-AV-01" in aviso,
                    f"{[(x['origen'], x['estado']) for x in t]} | {aviso[:60]}")
        aviso_vuelta = cambiar_fecha("2026-10-23", av02)
        r.comprobar("La indicación que salió de la lista vuelve con lo anotado (fecha corregida)",
                    pg.evaluate("state.acta.seguimiento.map(s => [s.id, s.estado, s.nuevoPlazo, s.comentario])") == anotado
                    and pg.evaluate("state.acta.seguimientoNotas.length") == 0 and "Seguimiento actualizado" in aviso_vuelta)
        aviso = cambiar_fecha("2026-10-01", None)          # antes de cualquier acta emitida: sin seguimiento
        guardadas = pg.evaluate("state.acta.seguimientoNotas.map(n => n.id)")
        cambiar_fecha("2026-10-23", av02)
        r.comprobar("Sin acta anterior la lista queda vacía, pero lo anotado se guarda y vuelve al corregir la fecha",
                    "sin seguimiento" in aviso and guardadas == [x[0] for x in anotado]
                    and pg.evaluate("state.acta.seguimiento.map(s => [s.id, s.estado, s.nuevoPlazo, s.comentario])") == anotado,
                    f"{aviso[:70]} | {guardadas}")
        pg.evaluate("irAFase(7)")                         # queda en Revisión, como si se fuera a emitir más tarde
        pg.evaluate("guardarYa()")

        # ══ «Corregir» la AV-02: la Rev. 1 conserva condiciones y seguimiento, que no se recarga ══
        abrir_mis_actas(pg)
        pg.click(f".acta-card[data-acta='{av02}'] button:has-text('Corregir')")
        pg.wait_for_selector("#promptOverlay.open")
        pg.fill("#promptTexto", "Se completa el comentario de la indicación AV-01 · 1")
        pg.click("#promptAceptar")
        esperar_aviso(pg)
        rev1 = pg.evaluate("state.acta.id")
        original, copia = leer(pg, av02), leer(pg, rev1)
        r.comprobar("«Corregir»: la Rev. 1 copia condiciones, seguimiento y su origen",
                    copia["condiciones"] == original["condiciones"] == {"meteo": ["lluvia"], "temperatura": "9"}
                    and copia["seguimiento"] == original["seguimiento"] and copia["seguimientoOrigen"] == original["seguimientoOrigen"])
        avanzar_hasta(pg, 5)
        pg.evaluate("cargarSeguimiento()")
        r.comprobar("En la revisión no se recarga el seguimiento: muestra lo anotado en la Rev. 0",
                    [x["estado"] for x in pg.evaluate(TARJETAS)] == ["cumplida", "pendiente", "anulada"])
        pg.fill("#segCom0", "Reparación ejecutada y comprobada en obra")
        pg.click("#btnNext")
        firmar_todos(cdp, pg)
        pg.click("#btnNext")
        pg.click("#btnNext")
        ruta_r1 = SALIDA / "S-AV02-Rev1.pdf"
        descargar(pg, "#btnDescargar", ruta_r1)
        pg.wait_for_function("state.acta.estado === 'emitida'")
        texto_r1 = texto_pdf(ruta_r1)
        r.comprobar("PDF de la Rev. 1: seguimiento con el comentario corregido",
                    "4 · SEGUIMIENTO DE INDICACIONES ANTERIORES" in texto_r1 and "comprobada en obra" in texto_r1
                    and f"Nº {EXP}-AV-02 · Rev. 1" in texto_r1)

        # ── La AV-03 se abrió en Revisión: como ahora la última emitida es la Rev. 1, se actualiza y se avisa ──
        pg.evaluate("id => abrirActaGuardada(id)", av03)
        aviso = esperar_aviso(pg)
        origenes = pg.evaluate("state.acta.seguimiento.map(s => [origenCorto(s.origen, state.acta.datos.expediente), s.estado])")
        r.comprobar("Revisión: si se ha emitido otra acta del expediente, el seguimiento se actualiza y se avisa",
                    "Seguimiento actualizado" in aviso and pg.evaluate("state.fase") == 7
                    and pg.evaluate("state.acta.seguimientoOrigen.actaId") == rev1
                    and origenes == [["AV-01 · 2", "cumplida"], ["AV-02 Rev. 1 · 1", "pendiente"]], f"{aviso[:60]} | {origenes}")

        # ══ Registro de actas en Excel (mejora 3) ══
        abrir_mis_actas(pg)
        pg.screenshot(path=str(SALIDA / "s-mis-actas.png"))
        nombre = descargar(pg, "#btnRegistro", SALIDA / "S-registro.xlsx")
        aviso = esperar_aviso(pg)
        r.comprobar("Registro: .xlsx en Descargas con el nombre y la fecha",
                    re.fullmatch(r"Registro_actas_TBC_\d{4}-\d{2}-\d{2}_\d{4}\.xlsx", nombre) is not None
                    and "Descargas" in aviso and "4 actas y 4 indicaciones" in aviso, f"{nombre} · {aviso[:80]}")
        with zipfile.ZipFile(SALIDA / "S-registro.xlsx") as z:
            integro = z.testzip() is None and z.namelist()[0] == "[Content_Types].xml"
        libro = load_workbook(SALIDA / "S-registro.xlsx")
        r.comprobar("Registro: ZIP íntegro con las hojas Actas e Indicaciones",
                    integro and libro.sheetnames == ["Actas", "Indicaciones"], str(libro.sheetnames))
        hoja = libro["Actas"]
        filas = list(hoja.iter_rows(min_row=2, values_only=True))
        r.comprobar("Hoja Actas: todas las actas del teléfono con su estado",
                    [(f[0], f[1], f[2]) for f in filas] == [(f"{EXP}-AV-01", 0, "Emitida"), (f"{EXP}-AV-02", 0, "Sustituida (Rev. 1)"),
                                                            (f"{EXP}-AV-02", 1, "Emitida"), (f"{EXP}-AV-03", 0, "Borrador")],
                    str([(f[0], f[1], f[2]) for f in filas]))
        r.comprobar("Hoja Actas: fechas reales (dd/mm/aaaa y dd/mm/aaaa hh:mm), condiciones, ubicación y asistentes",
                    isinstance(filas[0][3], datetime.datetime) and filas[0][3].date() == datetime.date(2026, 10, 8)
                    and hoja["D2"].number_format == "dd/mm/yyyy" and isinstance(filas[0][18], datetime.datetime)
                    and hoja["S2"].number_format == "dd/mm/yyyy hh:mm" and filas[3][18] is None
                    and filas[0][10] == "Nuboso, viento · 14,5 °C" and filas[0][11] == "37.389210, -5.984460 (±8 m)"
                    and filas[0][12] == "Juan Pérez Gómez (Constructora Sur · Jefe de obra)", str(filas[0][:13]))
        r.comprobar("Hoja Actas: objeto (tipo) y detalle en columnas propias, técnico y su cargo",
                    filas[0][6:10] == ("Inspección de Tajos de Obra", "Revisión de la estructura de planta primera", TECNICO,
                                       "Director de Obra"), str(filas[0][6:10]))
        r.comprobar("Hoja Actas: fotos, indicaciones y recuento del seguimiento",
                    [f[13:18] for f in filas] == [(1, 3, None, None, None), (1, 1, 1, 1, 1), (1, 1, 1, 1, 1), (1, 0, 1, 1, 0)],
                    str([f[13:18] for f in filas]))
        r.comprobar("Hoja Actas: fila 1 fija, autofiltro, cabecera destacada y estado en color",
                    hoja.freeze_panes == "A2" and hoja.auto_filter.ref == "A1:S5" and hoja["A1"].font.b
                    and hoja["A1"].fill.fgColor.rgb == "FFE8E8E8" and hoja["C2"].font.color.rgb == "FF1E7A3D",
                    f"{hoja.freeze_panes} · {hoja.auto_filter.ref}")
        hoja = libro["Indicaciones"]
        filas = list(hoja.iter_rows(min_row=2, values_only=True))
        vence = lambda plazo: "Sí" if plazo < hoy else "No"   # noqa: E731
        esperadas = [
            ("AV-01", 1, "Cumplida", "No", "AV-02 Rev. 1", "Reparación ejecutada y comprobada en obra"),
            ("AV-01", 2, "Pendiente", vence("2026-11-05"), "AV-02 Rev. 1", "Se reitera a la propiedad"),
            ("AV-01", 3, "Anulada", "No", "AV-02 Rev. 1", MOTIVO),
            ("AV-02 Rev. 1", 1, "Pendiente", vence("2026-10-20"), None, None),
        ]
        r.comprobar("Hoja Indicaciones: las de las actas emitidas vigentes, con el estado del último seguimiento",
                    [(f[1], f[2], f[8], f[9], f[10], f[11]) for f in filas] == esperadas,
                    str([(f[1], f[2], f[8], f[9], f[10]) for f in filas]))
        fechas = [(f[3], f[6], f[7]) for f in filas]
        dia = lambda v: v.date().isoformat() if isinstance(v, datetime.datetime) else v   # noqa: E731
        r.comprobar("Hoja Indicaciones: fecha de la visita, plazo vigente y plazo inicial como fechas",
                    [tuple(dia(v) for v in x) for x in fechas] == [("2026-10-08", "2026-10-12", "2026-10-12"),
                                                                    ("2026-10-08", "2026-11-05", "2026-10-30"),
                                                                    ("2026-10-08", None, None),
                                                                    ("2026-10-15", "2026-10-20", "2026-10-20")]
                    and hoja["G2"].number_format == "dd/mm/yyyy", str(fechas[:2]))
        r.comprobar("Hoja Indicaciones: expediente, texto y responsable; fila 1 fija y autofiltro",
                    all(f[0] == EXP for f in filas) and filas[0][4] == INDICACIONES_AV01[0][0]
                    and filas[1][5] == "Propiedad · Promotora" and hoja.freeze_panes == "A2" and hoja.auto_filter.ref == "A1:L5")

        # ══ Exportar e importar con seguimiento ══
        descargar(pg, f".acta-card[data-acta='{rev1}'] button:has-text('Exportar')", SALIDA / "S-rev1.zip")
        esperar_aviso(pg)
        descargar(pg, f".acta-card[data-acta='{av03}'] button:has-text('Exportar')", SALIDA / "S-av03.zip")
        esperar_aviso(pg)
        pg.evaluate("cerrarMisActas()")
        with zipfile.ZipFile(SALIDA / "S-rev1.zip") as z:
            carpeta = next(n for n in z.namelist() if n.endswith("acta.json")).rsplit("/", 1)[0]
            ficha = json.loads(z.read(f"{carpeta}/acta.json"))
            foto = z.read(f"{carpeta}/fotos/01.jpg")
        r.comprobar("Exportación: acta.json con el esquema nuevo, las condiciones y el seguimiento",
                    ficha["app"] == version and ficha["acta"]["esquema"] == esquema and len(ficha["acta"]["seguimiento"]) == 3
                    and ficha["acta"]["condiciones"] == {"meteo": ["lluvia"], "temperatura": "9"}
                    and ficha["acta"]["seguimientoOrigen"]["actaId"] == av01)
        manipulada = copy.deepcopy(ficha)
        seg = manipulada["acta"]["seguimiento"]
        seg[1].update(estado="hecha", nuevoPlazo="mañana", comentario="x" * 1000, plazo="05-11-2026", plazoInicial="2026-02-31")
        manipulada["acta"].update(
            id="acta-manipulada-1", seguimiento=[seg[0], seg[0], seg[1], {"id": "../../x", "texto": "sin id válido"}, "texto"],
            condiciones={"meteo": ["lluvia", "tornado", 5, "nuboso"], "temperatura": "99"},
            seguimientoOrigen={"actaId": "<img src=x>", "acta": "AV"}, versionEmision="<b>3</b>",
            seguimientoNotas=[{"id": "ind-nota", "estado": "anulada", "nuevoPlazo": "2026-12-01", "comentario": "y" * 400},
                              {"id": seg[0]["id"], "estado": "cumplida"}, {"id": "ind-nota"}, {"id": "a/b"}])
        with zipfile.ZipFile(SALIDA / "S-manipulada.zip", "w") as z:
            z.writestr("M/acta.json", json.dumps(manipulada))
            z.writestr("M/fotos/01.jpg", foto)

        ctx_b = contexto_movil(p, navegador)
        vigilar_red(ctx_b, externas)
        pb = abrir(ctx_b, errores)
        # Un borrador traído a un teléfono que no tiene su acta de origen conserva su seguimiento.
        pb.set_input_files("#importInput", str(SALIDA / "S-av03.zip"))
        esperar_aviso(pb)
        pb.evaluate("id => abrirActaGuardada(id)", av03)
        pb.wait_for_function(f"state.acta.id === '{av03}'")
        pb.evaluate("irAFase(5)")
        pb.evaluate("cargarSeguimiento()")
        traido, original = leer(pb, av03), leer(pg, av03)
        r.comprobar("Un borrador traído de otro teléfono, sin su acta de origen, conserva el seguimiento (sin avisos)",
                    traido["seguimiento"] == original["seguimiento"] and len(traido["seguimiento"]) == 2
                    and traido["seguimientoOrigen"] == original["seguimientoOrigen"]
                    and not pb.evaluate("document.getElementById('alertOverlay').classList.contains('open')"),
                    str([s["estado"] for s in traido["seguimiento"]]))
        # …y si en ese teléfono se pone mal la fecha (antes de su origen) y se corrige, la lista vuelve entera.
        pb.evaluate("irAFase(1)")
        pb.fill("#fechaVisita", "2026-10-14")
        pb.evaluate("irAFase(5)")
        pb.wait_for_function("state.acta.seguimientoOrigen === null")
        sin_lista = [esperar_aviso(pb), pb.evaluate("[state.acta.seguimiento.length, Boolean(state.acta.seguimientoGuardado)]")]
        pb.evaluate("irAFase(1)")
        pb.fill("#fechaVisita", "2026-10-23")
        pb.evaluate("irAFase(5)")
        pb.wait_for_function("state.acta.seguimientoOrigen !== null")
        esperar_aviso(pb)
        vuelta = pb.evaluate("[state.acta.seguimiento, state.acta.seguimientoOrigen, state.acta.seguimientoGuardado]")
        r.comprobar("Sin su acta de origen en el teléfono, una fecha mal puesta y corregida no pierde la lista",
                    "sin seguimiento" in sin_lista[0] and sin_lista[1] == [0, True] and vuelta[0] == original["seguimiento"]
                    and vuelta[1] == original["seguimientoOrigen"] and vuelta[2] is None, str(sin_lista[1]))
        pb.set_input_files("#importInput", str(SALIDA / "S-rev1.zip"))
        aviso = esperar_aviso(pb)
        en_a, en_b = leer(pg, rev1), leer(pb, rev1)
        r.comprobar("Importación en otro teléfono: se conservan condiciones, seguimiento, su origen y la versión de emisión",
                    "1 acta(s) nueva(s)" in aviso
                    and all(en_b[k] == en_a[k] for k in ("condiciones", "seguimiento", "seguimientoOrigen", "versionEmision")),
                    aviso)
        pb.evaluate("id => abrirActaGuardada(id)", rev1)
        pb.wait_for_function("state.fase === 8")
        pb.wait_for_function("document.getElementById('emitPdfEstado').textContent.startsWith('PDF listo')", timeout=60000)
        descargar(pb, "#btnDescargar", SALIDA / "S-AV02-Rev1-B.pdf")
        r.comprobar("El PDF del acta importada es idéntico al original", texto_pdf(SALIDA / "S-AV02-Rev1-B.pdf") == texto_r1)
        pb.set_input_files("#importInput", str(SALIDA / "S-manipulada.zip"))
        aviso = esperar_aviso(pb)
        m = leer(pb, "acta-manipulada-1") or {}
        seg_m = m.get("seguimiento", [])
        r.comprobar("Seguimiento manipulado: sin repetidas ni ids no válidos, estado conocido, fechas reales y longitudes acotadas",
                    "1 acta(s) nueva(s)" in aviso and [s["id"] for s in seg_m] == [seg[0]["id"], seg[1]["id"]]
                    and seg_m[1]["estado"] == "pendiente" and seg_m[1]["nuevoPlazo"] == "" and seg_m[1]["plazo"] == ""
                    and seg_m[1]["plazoInicial"] == "" and len(seg_m[1]["comentario"]) == 300, aviso)
        r.comprobar("Condiciones, origen, notas y versión manipulados: solo lo válido",
                    m.get("condiciones") == {"meteo": ["nuboso", "lluvia"], "temperatura": ""} and m.get("seguimientoOrigen") is None
                    and m.get("seguimientoNotas") == [{"id": "ind-nota", "estado": "anulada", "nuevoPlazo": "", "comentario": "y" * 300}]
                    and "versionEmision" not in m, f"{m.get('condiciones')} | {m.get('seguimientoNotas', [{}])[:1]}")
        ctx_b.close()

        # ══ Migración de las actas de la v2.3 (esquema 2) ══
        ctx_c = contexto_movil(p, navegador)
        vigilar_red(ctx_c, externas)
        pc = abrir(ctx_c, errores)
        pc.evaluate("""async () => {
            const imagen = await (await fetch('tests/salida/fotos/foto_horizontal.jpg')).blob();
            const base = { esquema: 2, creada: '2026-10-01T08:00:00.000Z', modificada: '2026-10-01T09:00:00.000Z', revision: 0,
              gps: { lat: '37.389210', lng: '-5.984460', acc: 8, hora: '2026-10-01T08:05:00.000Z' },
              sinGps: { activo: false, referencia: '', lat: '', lng: '' }, asistentes: [], observaciones: 'Acta de la v2.3',
              firma: [{ refW: 300, refH: 150, puntos: [[10, 10], [60, 40], [120, 20], [180, 60], [240, 30]] }], firmasAsistentes: {} };
            const datos = { numAuto: false, tecnico: 'Blas González González', cargoTecnico: '', expediente: 'EXP-V23',
              descripcion: 'Visita registrada con la v2.3', hora: '09:00' };
            await Almacen.guardarActa({ ...base, id: 'acta-v23-emitida', estado: 'emitida', emitida: '2026-10-01T10:00:00.000Z',
              fase: 8, revisiones: [{ rev: 0, fecha: '2026-10-01T10:00:00.000Z', motivo: 'Emisión inicial' }],
              datos: { ...datos, actaNum: 'EXP-V23-AV-01', fecha: '2026-10-01' },
              indicaciones: [{ id: 'ind-v23-1', texto: 'Indicación de la v2.3', responsable: 'Constructora Sur', plazo: '2026-10-05' }],
              fotos: [{ id: 'foto-v23-1', ancho: 4032, alto: 3024, comentario: 'Foto de la v2.3', hora: '2026-10-01 09:10',
                        horaFuente: 'movil', gps: null }] });
            await Almacen.guardarFoto('foto-v23-1', 'acta-v23-emitida', imagen);
            await Almacen.guardarActa({ ...base, id: 'acta-v23-borrador', estado: 'borrador', emitida: null, fase: 5,
              revisiones: [{ rev: 0, fecha: null, motivo: 'Emisión inicial' }], indicaciones: [], fotos: [],
              datos: { ...datos, actaNum: 'EXP-V23-AV-02', fecha: '2026-10-10' } });
            localStorage.setItem('tbc-acta-actual', 'acta-v23-borrador');
        }""")
        pc.reload()
        pc.wait_for_function("state.acta && state.acta.id === 'acta-v23-borrador'")
        pc.evaluate("cargarSeguimiento()")
        migrada = pc.evaluate("""({ fase: state.fase, esquema: state.acta.esquema, condiciones: state.acta.condiciones,
            seguimiento: state.acta.seguimiento.map(s => [s.id, s.estado, origenCorto(s.origen, state.acta.datos.expediente), s.plazo]) })""")
        r.comprobar("Un borrador de la v2.3 se abre migrado y en su misma fase (sin el cambio de fases de la v2.1)",
                    migrada["fase"] == 5 and migrada["esquema"] == esquema
                    and migrada["condiciones"] == {"meteo": [], "temperatura": ""}, str(migrada))
        r.comprobar("…y hace el seguimiento de las indicaciones de un acta emitida con la v2.3",
                    migrada["seguimiento"] == [["ind-v23-1", "pendiente", "AV-01 · 1", "2026-10-05"]], str(migrada["seguimiento"]))
        pc.evaluate("abrirActaGuardada('acta-v23-emitida')")
        pc.wait_for_function("state.acta.id === 'acta-v23-emitida' && state.fase === 8")
        pc.wait_for_function("document.getElementById('emitPdfEstado').textContent.startsWith('PDF listo')", timeout=60000)
        descargar(pc, "#btnDescargar", SALIDA / "S-v23-migrada.pdf")
        texto_v23 = texto_pdf(SALIDA / "S-v23-migrada.pdf")
        r.comprobar("El acta emitida con la v2.3 rehace el PDF que se distribuyó: secciones de la v2.3, sin enlaces al mapa",
                    all(s in texto_v23 for s in SECCIONES_V23) and "SEGUIMIENTO" not in texto_v23 and "CONDICIONES" not in texto_v23
                    and "Ver en el mapa" not in texto_v23 and "Emitida 01/10/2026" in texto_v23
                    and enlaces(SALIDA / "S-v23-migrada.pdf") == [], ", ".join(s for s in SECCIONES_V23 if s not in texto_v23))
        marco_v23 = marco_pdf(SALIDA / "S-v23-migrada.pdf")
        r.comprobar("…con el formato de entonces: «Acta de Visita de Inspección», cabecera y pie en negro",
                    "ACTA DE VISITA DE INSPECCIÓN" in texto_v23 and "Acta de Campo" not in texto_v23
                    and marco_v23 == [MARCO_NEGRO] * len(marco_v23), str(marco_v23))

        # ── Un borrador que se abre directamente en Emisión comprueba el seguimiento antes de preparar el PDF ──
        pc.evaluate("""async () => {
            const posterior = await Almacen.leerActa('acta-v23-emitida');
            Object.assign(posterior, { id: 'acta-v23-posterior', emitida: '2026-10-05T10:00:00.000Z', modificada: '2026-10-05T10:00:00.000Z',
              fotos: [], indicaciones: [{ id: 'ind-v23-2', texto: 'Indicación posterior', responsable: 'Constructora Sur', plazo: '2026-10-20' }] });
            posterior.datos = { ...posterior.datos, actaNum: 'EXP-V23-AV-01B', fecha: '2026-10-05' };
            await Almacen.guardarActa(posterior);
            const borrador = await Almacen.leerActa('acta-v23-borrador');
            borrador.fase = 8;
            await Almacen.guardarActa(borrador);
        }""")
        pc.evaluate("abrirActaGuardada('acta-v23-borrador')")
        aviso = esperar_aviso(pc)
        r.comprobar("Si el seguimiento ha cambiado, no se prepara el PDF: vuelve a Acuerdos y avisa",
                    "Seguimiento actualizado" in aviso and "EXP-V23-AV-01B" in aviso and pc.evaluate("state.fase") == 5
                    and pc.evaluate("state.acta.seguimiento.map(s => s.id)") == ["ind-v23-2"]
                    and pc.evaluate("pdfEmision.promesa === null || pdfEmision.clave !== clavePDF(state.acta)"), aviso[:90])

        # ── Borrador esperando en Emisión con el PDF preparado: llega importada un acta emitida posterior ──
        pc.evaluate("irAFase(8)")
        pc.wait_for_function("document.getElementById('emitPdfEstado').textContent.startsWith('PDF listo')", timeout=60000)
        intermedia = {"formato": "tbc-acta", "version": 1, "app": version, "exportada": "2026-10-23T08:00:00.000Z", "acta": {
            "id": "acta-v23-intermedia", "esquema": esquema, "estado": "emitida", "emitida": "2026-10-08T10:00:00.000Z",
            "versionEmision": version, "creada": "2026-10-08T08:00:00.000Z", "modificada": "2026-10-08T10:00:00.000Z",
            "fase": 8, "revision": 0, "revisiones": [{"rev": 0, "fecha": "2026-10-08T10:00:00.000Z", "motivo": "Emisión inicial"}],
            "datos": {"actaNum": "EXP-V23-AV-01C", "tecnico": "Otro técnico", "cargoTecnico": "", "expediente": "EXP-V23",
                      "descripcion": "Visita intermedia", "fecha": "2026-10-08", "hora": "09:00"},
            "gps": {"lat": "37.389210", "lng": "-5.984460", "acc": 8, "hora": None},
            "sinGps": {"activo": False, "referencia": "", "lat": "", "lng": ""}, "asistentes": [], "observaciones": "",
            "indicaciones": [{"id": "ind-v23-3", "texto": "Indicación de la visita intermedia", "responsable": "Constructora Sur",
                              "plazo": "2026-10-25"}],
            "fotos": [{"archivo": "fotos/01.jpg", "ancho": 1200, "alto": 1600, "comentario": "", "hora": "2026-10-08 09:10",
                       "horaFuente": "movil", "gps": None}],
            "firma": [{"refW": 300, "refH": 150, "puntos": [[10, 10], [60, 40], [120, 20], [180, 60], [240, 30]]}],
            "firmasAsistentes": {}}}
        with zipfile.ZipFile(SALIDA / "S-intermedia.zip", "w") as z:
            z.writestr("I/acta.json", json.dumps(intermedia))
            z.writestr("I/fotos/01.jpg", foto)
        pc.set_input_files("#importInput", str(SALIDA / "S-intermedia.zip"))
        aviso = esperar_aviso(pc)
        r.comprobar("Con el PDF ya preparado, si al importar llega un acta emitida posterior, no se emite: vuelve a Acuerdos",
                    "1 acta(s) nueva(s)" in aviso and "seguimiento de indicaciones del acta en pantalla ha cambiado" in aviso
                    and pc.evaluate("state.fase") == 5 and pc.evaluate("state.acta.seguimientoOrigen.actaId") == "acta-v23-intermedia"
                    and pc.evaluate("state.acta.seguimiento.map(s => s.id)") == ["ind-v23-3"], aviso[:120])

        # ── Aunque el PDF ya esté preparado, «Compartir» y «Descargar» esperan a que termine la comprobación ──
        pc.evaluate("irAFase(8)")
        pc.wait_for_function("document.getElementById('emitPdfEstado').textContent.startsWith('PDF listo')", timeout=60000)
        pc.evaluate("""(() => { const original = Almacen.listarActas; window.__listarActas = original;
            Almacen.listarActas = async () => { await new Promise(ok => setTimeout(ok, 1500)); return original(); }; })()""")
        pc.evaluate("irAFase(7)")
        pc.evaluate("irAFase(8)")
        durante = pc.evaluate("[btnDescargar.disabled, emitPdfEstado.textContent]")
        pc.wait_for_function("document.getElementById('emitPdfEstado').textContent.startsWith('PDF listo')", timeout=10000)
        despues = pc.evaluate("[btnDescargar.disabled, state.acta.estado]")
        pc.evaluate("Almacen.listarActas = window.__listarActas")
        r.comprobar("Con el PDF ya preparado, los botones esperan a que termine la comprobación del seguimiento",
                    durante[0] is True and "Comprobando el seguimiento" in durante[1] and despues == [False, "borrador"],
                    f"{durante} → {despues}")
        ctx_c.close()

        # ══ Service worker: precarga con «?v=<VERSION>» (mejora 7) ══
        ctx_d = contexto_movil(p, navegador)
        pd = abrir(ctx_d, errores)
        pd.wait_for_function("navigator.serviceWorker.controller !== null", timeout=20000)
        cachés = pd.evaluate("caches.keys()")
        urls = pd.evaluate(f"caches.open('acta-tbc-v{version}').then(c => c.keys()).then(ks => ks.map(k => k.url))")
        r.comprobar("Service worker: la caché de la versión guarda los 9 archivos pedidos con «?v=<VERSION>»",
                    cachés == [f"acta-tbc-v{version}"] and len(urls) == 9 and all(u.endswith(f"?v={version}") for u in urls),
                    f"{cachés} · {urls[:2]}")
        ctx_d.set_offline(True)
        servidas = pd.evaluate("Promise.all(['index.html', 'lib/fuentes-pdf.js', './'].map(u => fetch(u).then(r => r.ok, () => false)))")
        pd.reload()
        pd.wait_for_function("typeof state !== 'undefined' && state.acta !== null", timeout=15000)
        r.comprobar("Sin conexión, la app y sus archivos (pedidos sin «?v=») salen de la caché",
                    all(servidas) and pd.evaluate("VERSION") == version, str(servidas))
        ctx_d.close()

        r.comprobar("Sin peticiones a internet", not externas, "; ".join(externas[:3]))
        r.comprobar("Sin errores JS", not errores, "; ".join(errores[:3]))
        ctx.close()
        navegador.close()
    srv.shutdown()
    return r.resumen()


if __name__ == "__main__":
    sys.exit(main())
